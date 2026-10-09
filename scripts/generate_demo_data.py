#!/usr/bin/env python3
"""Generate the synthetic demo dataset for Gel.

Run:  scripts/.venv/bin/python scripts/generate_demo_data.py

Deletes and regenerates demo-data/ from scratch (deterministic, fixed seed).
ALL people, IDs, companies, banks and numbers are fictional.

Fonts: uses macOS system TrueType fonts (Arial, Times New Roman, Verdana,
Courier New) from /System/Library/Fonts/Supplemental so the peso sign
renders and the PDF text layer is real, extractable Unicode text.
"""
from __future__ import annotations

import io
import json
import re
import shutil
import sys
import zlib
from datetime import date, timedelta
from pathlib import Path
import random

from PIL import Image, ImageChops, ImageDraw, ImageFilter, ImageFont
from reportlab import rl_config
from reportlab.lib.pagesizes import A4
from reportlab.lib.utils import ImageReader
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.pdfgen import canvas as rl_canvas
import docx
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

SEED = 20261009
TODAY = date(2026, 10, 9)
ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / "demo-data"
HR = OUT / "HR Files"
FOOTER = "Synthetic demo data — fictional person"
PESO = "₱"
EN = "–"

rl_config.invariant = 1  # deterministic PDF bytes (no timestamps / random IDs)

# ---------------------------------------------------------------------------
# Fonts
# ---------------------------------------------------------------------------
FONT_DIR = Path("/System/Library/Fonts/Supplemental")
FONT_FILES = {
    "sans": "Arial.ttf", "sans-b": "Arial Bold.ttf", "sans-i": "Arial Italic.ttf",
    "serif": "Times New Roman.ttf", "serif-b": "Times New Roman Bold.ttf",
    "serif-i": "Times New Roman Italic.ttf",
    "verd": "Verdana.ttf", "verd-b": "Verdana Bold.ttf", "verd-i": "Verdana Italic.ttf",
    "mono": "Courier New.ttf", "mono-b": "Courier New Bold.ttf",
}
for _k, _f in FONT_FILES.items():
    _p = FONT_DIR / _f
    if not _p.exists():
        sys.exit(f"Missing font {_p}. This generator expects macOS system fonts.")
    pdfmetrics.registerFont(TTFont(_k, str(_p)))

_PIL_FONTS: dict = {}


def pil_font(key: str, px: float):
    k = (key, round(px, 1))
    if k not in _PIL_FONTS:
        _PIL_FONTS[k] = ImageFont.truetype(str(FONT_DIR / FONT_FILES[key]), round(px, 1))
    return _PIL_FONTS[k]


def sw(s: str, font: str, size: float) -> float:
    return pdfmetrics.stringWidth(s, font, size)


def wrap(s: str, font: str, size: float, maxw: float) -> list[str]:
    out = []
    for para in s.split("\n"):
        cur = ""
        for w in para.split(" "):
            t = w if not cur else cur + " " + w
            if not cur or sw(t, font, size) <= maxw:
                cur = t
            else:
                out.append(cur)
                cur = w
        out.append(cur)
    return out


BLACK = (0, 0, 0)
DARK = (0.15, 0.15, 0.17)
GRAY = (0.45, 0.45, 0.48)
LIGHT = (0.62, 0.62, 0.64)
RULE = (0.75, 0.75, 0.78)
TEAL = (0.05, 0.36, 0.40)
NAVY = (0.12, 0.20, 0.36)
BAND = (0.90, 0.92, 0.93)
WHITE = (1, 1, 1)


# ---------------------------------------------------------------------------
# Tiny page model: layout once, render to PDF (text layer) or to image (scans)
# ---------------------------------------------------------------------------
class Doc:
    W, H = A4

    def __init__(self, title, ml=54, mr=54, mt=54, mb=64, on_new_page=None, page_numbers=True):
        self.title = title
        self.ml, self.mr, self.mt, self.mb = ml, mr, mt, mb
        self.pages: list[list] = []
        self.pii: list[tuple[str, str]] = []
        self.on_new_page = on_new_page
        self.page_numbers = page_numbers
        self.y = mt
        self.new_page()

    @property
    def cw(self):
        return self.W - self.ml - self.mr

    def new_page(self):
        self.pages.append([])
        self.y = self.mt
        if self.on_new_page and len(self.pages) > 1:
            self.on_new_page(self)

    def ensure(self, h):
        if self.y + h > self.H - self.mb:
            self.new_page()

    # primitives (y is top-down; text y is the baseline)
    def text(self, x, y, s, font="sans", size=10, color=BLACK, anchor="l"):
        if anchor == "r":
            x -= sw(s, font, size)
        elif anchor == "c":
            x -= sw(s, font, size) / 2
        self.pages[-1].append(("text", x, y, s, font, size, color))

    def line(self, x1, y1, x2, y2, w=0.6, color=RULE):
        self.pages[-1].append(("line", x1, y1, x2, y2, w, color))

    def rect(self, x, y, w, h, fill=None, stroke=None, lw=0.6):
        self.pages[-1].append(("rect", x, y, w, h, fill, stroke, lw))

    # flow helpers
    def space(self, h):
        self.y += h

    def para(self, s, font="sans", size=10, x=None, w=None, leading=None, color=BLACK):
        x = self.ml if x is None else x
        w = (self.W - self.mr - x) if w is None else w
        leading = leading or size * 1.38
        for ln in wrap(s, font, size, w):
            self.ensure(leading)
            self.text(x, self.y + size, ln, font, size, color)
            self.y += leading

    def bullet(self, s, font="sans", size=10, x=None, indent=12, color=BLACK, sym="•"):
        x = self.ml if x is None else x
        leading = size * 1.38
        self.ensure(leading)
        self.text(x + 2, self.y + size, sym, font, size, color)
        self.para(s, font, size, x=x + indent, leading=leading, color=color)

    def mark(self, typ, value):
        self.pii.append((typ, value))
        return value

    def finalize(self):
        n = len(self.pages)
        for i, page in enumerate(self.pages):
            yb = self.H - 26
            page.append(("text", self.W / 2 - sw(FOOTER, "sans", 6.5) / 2, yb, FOOTER, "sans", 6.5, LIGHT))
            if self.page_numbers and n > 1:
                s = f"Page {i + 1} of {n}"
                page.append(("text", self.W - self.mr - sw(s, "sans", 7), yb, s, "sans", 7, LIGHT))

    def page_texts(self):
        return [re.sub(r"\s+", " ", " ".join(op[3] for op in p if op[0] == "text")) for p in self.pages]


def render_pdf(doc: Doc, path: Path):
    path.parent.mkdir(parents=True, exist_ok=True)
    c = rl_canvas.Canvas(str(path), pagesize=A4, invariant=1, pageCompression=1)
    c.setTitle(doc.title)
    c.setAuthor("Gel demo data generator")
    c.setSubject(FOOTER)
    H = doc.H
    for page in doc.pages:
        for op in page:
            if op[0] == "text":
                _, x, y, s, font, size, color = op
                c.setFont(font, size)
                c.setFillColorRGB(*color)
                c.drawString(x, H - y, s)
            elif op[0] == "line":
                _, x1, y1, x2, y2, w, color = op
                c.setLineWidth(w)
                c.setStrokeColorRGB(*color)
                c.line(x1, H - y1, x2, H - y2)
            elif op[0] == "rect":
                _, x, y, w, h, fill, stroke, lw = op
                if fill:
                    c.setFillColorRGB(*fill)
                if stroke:
                    c.setStrokeColorRGB(*stroke)
                    c.setLineWidth(lw)
                c.rect(x, H - y - h, w, h, fill=1 if fill else 0, stroke=1 if stroke else 0)
        c.showPage()
    c.save()


def _c255(c):
    return tuple(int(round(v * 255)) for v in c)


def render_page_image(page, dpi=170) -> Image.Image:
    s = dpi / 72.0
    img = Image.new("RGB", (int(A4[0] * s), int(A4[1] * s)), "white")
    d = ImageDraw.Draw(img)
    for op in page:
        if op[0] == "rect":
            _, x, y, w, h, fill, stroke, lw = op
            d.rectangle([x * s, y * s, (x + w) * s, (y + h) * s],
                        fill=_c255(fill) if fill else None,
                        outline=_c255(stroke) if stroke else None,
                        width=max(1, round(lw * s)) if stroke else 0)
    for op in page:
        if op[0] == "line":
            _, x1, y1, x2, y2, w, color = op
            d.line([(x1 * s, y1 * s), (x2 * s, y2 * s)], fill=_c255(color), width=max(1, round(w * s)))
        elif op[0] == "text":
            _, x, y, txt, font, size, color = op
            d.text((x * s, y * s), txt, font=pil_font(font, size * s), fill=_c255(color), anchor="ls")
    return img


# ---------------------------------------------------------------------------
# Synthetic identity generators
# ---------------------------------------------------------------------------
R = random.Random(SEED)

FEMALE = ["Maria", "Ana", "Kristine", "Jasmine", "Angelica", "Patricia", "Rowena", "Marites", "Lorna",
          "Jennifer", "Mary Grace", "Rachelle", "Joanna", "Liza", "Sheila", "Michelle", "Bea", "Andrea",
          "Princess", "Rosalinda", "Charisse", "Nicole", "Erlinda", "Jonalyn", "Hazel", "Aileen", "Divina",
          "Cherry Mae", "Rhea", "Katrina"]
MALE = ["Jose", "Juan Carlo", "Mark Anthony", "John Paul", "Ramon", "Christian", "Jerome", "Rafael",
        "Paolo", "Miguel", "Arnel", "Ronaldo", "Jayson", "Kenneth", "Noel", "Dennis", "Rommel", "Vincent",
        "Bryan", "Jericho", "Emmanuel", "Francis", "Alvin", "Reynaldo", "Gilbert", "Carlo", "Joel"]
# Reyes / Santos / Cruz (and Dela Cruz) are deliberately excluded: only the 3 demo applicants use them.
SURNAMES = ["Garcia", "Mendoza", "Bautista", "Villanueva", "Ramos", "Castillo", "Flores", "Gonzales",
            "Torres", "Navarro", "Salazar", "Pascual", "Mercado", "Manalo", "Del Rosario", "Soriano",
            "Lopez", "Fernandez", "Valdez", "Rivera", "Gutierrez", "Tolentino", "Panganiban", "Macaraeg",
            "Dimaculangan", "Sison", "Magbanua", "Evangelista", "Javier", "Marquez", "Dizon", "Hernandez",
            "Perez", "Morales", "Cortez", "Ignacio", "Galang", "Samonte", "Buenaventura", "De Guzman",
            "Alcantara", "Cabrera", "Velasco", "Bernardo", "Saldana", "Lim", "Tan", "Quiambao", "Dimalanta",
            "Ferrer", "Abad", "Mariano", "Robles", "Umali", "Vergara", "Yap", "Zamora", "Aguilar", "Domingo"]

# (city, province, [(barangay, zip)])
LOCALES = [
    ("Quezon City", "Metro Manila", [("Bagong Pag-asa", "1105"), ("Commonwealth", "1121"), ("Batasan Hills", "1126"),
                                     ("Kamuning", "1103"), ("Tandang Sora", "1116")]),
    ("Makati City", "Metro Manila", [("Poblacion", "1210"), ("Pio del Pilar", "1230"), ("Bangkal", "1233")]),
    ("Pasig City", "Metro Manila", [("Kapitolyo", "1603"), ("Pinagbuhatan", "1602"), ("Rosario", "1609")]),
    ("Mandaluyong City", "Metro Manila", [("Highway Hills", "1550"), ("Plainview", "1550"), ("Barangka Ilaya", "1550")]),
    ("Taguig City", "Metro Manila", [("Western Bicutan", "1630"), ("Ususan", "1639"), ("Lower Bicutan", "1632")]),
    ("Caloocan City", "Metro Manila", [("Bagong Silang", "1428"), ("Camarin", "1422")]),
    ("Antipolo City", "Rizal", [("San Roque", "1870"), ("Dela Paz", "1870"), ("Mayamot", "1870")]),
    ("Cainta", "Rizal", [("San Andres", "1900"), ("Santo Domingo", "1900")]),
    ("Bacoor City", "Cavite", [("Molino III", "4102"), ("Niog", "4102"), ("Talaba", "4102")]),
    ("Dasmariñas City", "Cavite", [("Salitran", "4114"), ("Sampaloc I", "4114"), ("Burol", "4114")]),
    ("Calamba City", "Laguna", [("Parian", "4027"), ("Real", "4027"), ("Canlubang", "4028")]),
    ("Santa Rosa City", "Laguna", [("Balibago", "4026"), ("Tagapo", "4026")]),
    ("City of San Fernando", "Pampanga", [("Dolores", "2000"), ("Sindalan", "2000")]),
    ("Malolos City", "Bulacan", [("Longos", "3000"), ("Sumapang Matanda", "3000")]),
    ("Cebu City", "Cebu", [("Lahug", "6000"), ("Guadalupe", "6000"), ("Talamban", "6000")]),
    ("Mandaue City", "Cebu", [("Banilad", "6014"), ("Tipolo", "6014")]),
    ("Davao City", "Davao del Sur", [("Buhangin", "8000"), ("Matina Crossing", "8000")]),
    ("Iloilo City", "Iloilo", [("Jaro", "5000"), ("Mandurriao", "5000")]),
    ("Baguio City", "Benguet", [("Aurora Hill", "2600"), ("Irisan", "2600")]),
    ("Lipa City", "Batangas", [("Marawoy", "4217"), ("Sabang", "4217")]),
    ("Cagayan de Oro City", "Misamis Oriental", [("Carmen", "9000"), ("Lapasan", "9000")]),
]
STREETS = ["Rizal St.", "Mabini St.", "Bonifacio Ave.", "Luna St.", "Sampaguita St.", "Narra St.", "Acacia St.",
           "Ilang-Ilang St.", "Kamagong St.", "Molave St.", "Gumamela St.", "Del Pilar St.", "Burgos St.",
           "Mahogany Rd.", "Jasmine St.", "Dahlia St.", "Everlasting St.", "Magsaysay Ave.", "Aguinaldo St.",
           "Katipunan St.", "Banaba St.", "Waling-Waling St."]

GLOBE_SMART = ["0905", "0906", "0915", "0916", "0917", "0926", "0927", "0935", "0936", "0945", "0953", "0955",
               "0956", "0965", "0966", "0975", "0977", "0995", "0997", "0907", "0908", "0909", "0910", "0912",
               "0918", "0919", "0920", "0921", "0928", "0929", "0939", "0947", "0949", "0951", "0998", "0999",
               "0991", "0993"]
EMAIL_DOMAINS = ["example-mail.ph", "pinoymail.example", "inbox-demo.ph", "sulat.example"]

COMPANY = "Bayanihan Outsourcing Corp."
COMPANY_ADDR = "8/F Halimbawa Tower, 21 Example Ave., Ortigas Center, Pasig City, Metro Manila 1605"
PAYROLL_BANK = "Bangko Halimbawa"
MONTHS = ["January", "February", "March", "April", "May", "June", "July", "August", "September", "October",
          "November", "December"]


def digits(n, rng=R):
    return "".join(rng.choice("0123456789") for _ in range(n))


def gen_sss(rng=R):
    return f"{rng.randint(1, 34):02d}-{digits(7, rng)}-{digits(1, rng)}"


def gen_tin(rng=R):
    return f"{rng.randint(100, 499)}-{digits(3, rng)}-{digits(3, rng)}-" + ("00000" if rng.random() < 0.2 else "000")


def gen_philhealth(rng=R):
    return f"{rng.choice(['01', '02', '03', '08', '12', '19'])}-{digits(9, rng)}-{digits(1, rng)}"


def gen_pagibig(rng=R):
    return f"1{digits(3, rng)}-{digits(4, rng)}-{digits(4, rng)}"


def gen_psn(rng=R):
    return f"{rng.randint(1000, 9999)}-{digits(4, rng)}-{digits(4, rng)}-{digits(4, rng)}"


def gen_passport(rng=R):
    return f"P{digits(7, rng)}{rng.choice('ABCDEFGHJKLMNPRSTUVWXYZ')}"


def gen_dl(rng=R):
    return f"{rng.choice('NDACBEFGHJK')}{rng.randint(1, 16):02d}-{rng.randint(10, 25):02d}-{digits(6, rng)}"


def gen_bank_account(rng=R):
    return f"{rng.randint(100, 999)}-{digits(4, rng)}-{digits(rng.choice([3, 4, 5]), rng)}"


def luhn_ok(num: str) -> bool:
    ds = [int(c) for c in num if c.isdigit()]
    tot = 0
    for i, d in enumerate(reversed(ds)):
        if i % 2 == 1:
            d *= 2
            if d > 9:
                d -= 9
        tot += d
    return tot % 10 == 0


def gen_card(prefix="4000", rng=R):
    body = prefix + digits(15 - len(prefix), rng)
    for chk in "0123456789":
        if luhn_ok(body + chk):
            n = body + chk
            return " ".join(n[i:i + 4] for i in range(0, 16, 4))
    raise RuntimeError


def gen_phone(intl=False, rng=R):
    p = rng.choice(GLOBE_SMART)
    a, b = digits(3, rng), digits(4, rng)
    return f"+63 {p[1:]} {a} {b}" if intl else f"{p} {a} {b}"


def gen_address(rng=R, locale=None):
    city, prov, brgys = locale or rng.choice(LOCALES)
    brgy, zipc = rng.choice(brgys)
    st = rng.choice(STREETS)
    f = rng.random()
    if f < 0.45:
        street = f"{rng.randint(3, 289)} {st}"
    elif f < 0.85:
        street = f"Blk {rng.randint(1, 48)} Lot {rng.randint(1, 30)}, {st}"
    else:
        street = f"Unit {rng.randint(2, 18)}{rng.choice('ABCDEF')}, {rng.choice(['Narra', 'Molave', 'Anahaw', 'Sampaguita'])} Residences, {rng.randint(10, 150)} {st}"
    return f"{street}, Brgy. {brgy}, {city}, {prov} {zipc}", city


def fmt_long(d: date) -> str:
    return f"{MONTHS[d.month - 1]} {d.day}, {d.year}"


def fmt_mdy(d: date) -> str:
    return f"{d.month:02d}/{d.day:02d}/{d.year}"


def rand_date(y0, y1, rng=R):
    start = date(y0, 1, 1)
    return start + timedelta(days=rng.randint(0, (date(y1, 12, 31) - start).days))


def age_on(dob: date, ref=TODAY):
    return ref.year - dob.year - ((ref.month, ref.day) < (dob.month, dob.day))


def peso(x: float) -> str:
    return f"{PESO}{x:,.2f}"


_used_surnames = {"Reyes", "Santos", "Cruz", "Villafuerte"}


def take_surname(rng=R):
    pool = [s for s in SURNAMES if s not in _used_surnames]
    s = rng.choice(pool)
    _used_surnames.add(s)
    return s


def make_person(rng=R, sex=None, first=None, middle=None, last=None, yob=(1984, 2001)):
    sex = sex or rng.choice("FM")
    first = first or rng.choice(FEMALE if sex == "F" else MALE)
    last = last or take_surname(rng)
    middle = middle or rng.choice([s for s in SURNAMES if s != last])
    dob = rand_date(*yob, rng=rng)
    addr, city = gen_address(rng)
    uname = f"{first.split()[0].lower()}.{re.sub(r'[^a-z]', '', last.lower())}"
    if rng.random() < 0.5:
        uname += str(rng.randint(1, 99))
    email = f"{uname}@{rng.choice(EMAIL_DOMAINS)}"
    return {
        "sex": sex, "first": first, "middle": middle, "last": last,
        "full": f"{first} {middle} {last}", "short": f"{first} {middle[0]}. {last}",
        "lfm": f"{last.upper()}, {first} {middle}",
        "dob": dob, "addr": addr, "city": city, "phone": gen_phone(rng.random() < 0.3, rng),
        "email": email, "sss": gen_sss(rng), "tin": gen_tin(rng), "philhealth": gen_philhealth(rng),
        "pagibig": gen_pagibig(rng), "civil": rng.choice(["Single", "Single", "Married", "Married", "Single"]),
        "pob": rng.choice(LOCALES)[0],
    }


# ---------------------------------------------------------------------------
# Ground truth bookkeeping
# ---------------------------------------------------------------------------
ENTITIES: list[dict] = []


def record(file_rel: str, pii: list[tuple[str, str]], page_texts: list[str]):
    seen = set()
    for typ, val in pii:
        found = [i + 1 for i, t in enumerate(page_texts) if val in t]
        if not found:
            raise RuntimeError(f"{file_rel}: planted {typ} {val!r} not found in rendered text")
        for p in found:
            key = (typ, val, p)
            if key not in seen:
                seen.add(key)
                ENTITIES.append({"file": file_rel, "type": typ, "value": val, "page": p})


def emit_pdf(doc: Doc, path: Path):
    doc.finalize()
    render_pdf(doc, path)
    record(str(path.relative_to(OUT)), doc.pii, doc.page_texts())
    return doc


# ---------------------------------------------------------------------------
# Resume content
# ---------------------------------------------------------------------------
PRIOR_COMPANIES = ["Kapitbahay Logistics Corp.", "Narra Ridge Foods Inc.", "Sampaguita Retail Group",
                   "Dagat-Bituin Manufacturing Corp.", "Luntiang Lambak Agri Corp.", "Halimbawa Holdings Inc.",
                   "Bituin Contact Solutions Inc.", "Anahaw Digital Ventures Inc.", "Tanglaw Accounting Partners",
                   "Perlas Halimbawa Hotel Group", "Kalayaan Example BPO Services", "Habagat Trading Co.",
                   "Mayumi Health Products Inc.", "Sinag Example Realty Corp.", "Talisay Demo Foods Corp."]
SCHOOLS = ["Pamantasang Bayan ng Halimbawa", "Colegio de San Ysidro Labrador", "Visayas Halimbawa State University",
           "Luzon Polytechnic Example College", "Mindanao Demo State University", "Kolehiyo ng Bagong Silang",
           "St. Lorenzo Example University", "Laguna Halimbawa Institute of Technology"]
HIGH_SCHOOLS = ["Halimbawa National High School", "San Ysidro Catholic High School", "Bagong Silang Integrated School",
                "Rizal Example National High School"]

ROLE_BULLETS = {
    "Payroll Specialist": [
        "Processed semi-monthly payroll for 1,200+ rank-and-file and supervisory employees across 4 sites",
        "Computed overtime, night differential, holiday pay and leave conversions based on approved timekeeping",
        "Prepared and remitted monthly SSS, PhilHealth and Pag-IBIG contributions and loan amortizations",
        "Computed withholding tax under the TRAIN law and prepared BIR Forms 1601-C and 2316",
        "Reconciled payroll registers with GL entries and resolved employee payroll disputes within 48 hours",
        "Handled final pay and 13th month pay computations for resigned and active employees"],
    "Payroll Officer": [
        "Owned end-to-end payroll processing for 650 employees using an HRIS payroll module",
        "Validated DTRs, overtime authorizations and leave balances before every payroll cut-off",
        "Prepared government remittance reports (SSS R3, PhilHealth RF-1, Pag-IBIG MCRF)",
        "Computed annualized withholding tax and year-end adjustments; issued BIR Form 2316",
        "Coordinated bank payroll enrollment and salary crediting with the payroll bank"],
    "Senior Payroll Analyst": [
        "Led a 4-person payroll team supporting 2,300 employees on semi-monthly payroll",
        "Audited payroll registers and variance reports prior to release; reduced payroll errors by 60%",
        "Implemented automated computation of night differential and holiday premiums in Excel/VBA",
        "Served as point person for SSS, PhilHealth, Pag-IBIG and BIR payroll compliance audits",
        "Prepared monthly payroll cost reports and accruals for Finance"],
    "Payroll Associate": [
        "Encoded timekeeping data and computed basic pay, overtime and statutory deductions",
        "Prepared payslips and payroll registers for review of the Payroll Supervisor",
        "Processed SSS salary loan and Pag-IBIG multi-purpose loan deductions",
        "Answered employee payroll inquiries via email and walk-in"],
    "Payroll Assistant": [
        "Assisted the Payroll Officer in encoding daily time records and overtime",
        "Prepared payroll summaries and printed payslips for distribution",
        "Filed payroll documents and government contribution reports"],
    "Timekeeping and Payroll Clerk": [
        "Consolidated biometric logs and validated tardiness, undertime and absences",
        "Prepared timekeeping reports used as basis for semi-monthly payroll",
        "Assisted in encoding payroll adjustments and leave applications"],
    "HR Operations Supervisor": [
        "Supervises a team of 5 HR associates handling employee records, onboarding and offboarding",
        "Maintains 201 files and HRIS data integrity for 1,500+ employees",
        "Implemented a digital onboarding checklist that cut pre-employment processing time by 30%",
        "Coordinates with Finance on compensation adjustments, benefits enrollment and clearance",
        "Prepares monthly HR operations dashboards for the HR Director"],
    "HR Generalist": [
        "Handles employee relations concerns, notices to explain and administrative hearings",
        "Administers HMO enrollment, leave management and employee engagement activities",
        "Supports recruitment for support functions and conducts onboarding orientation",
        "Maintains compliance with DOLE labor standards and company policies"],
    "HR Assistant": [
        "Maintained and updated 201 files, ID requests and certificates of employment",
        "Scheduled interviews, sent job offers and coordinated pre-employment requirements",
        "Prepared onboarding kits and conducted new hire orientation",
        "Encoded employee data in the HRIS and prepared monthly headcount reports"],
    "Recruitment Specialist": [
        "Managed full-cycle recruitment for 25-40 monthly openings in a high-volume BPO environment",
        "Sourced candidates via job portals, referrals and campus job fairs",
        "Conducted initial interviews, assessments and job offer discussions",
        "Maintained applicant tracking records and weekly pipeline reports"],
    "Recruitment Assistant": [
        "Screened resumes and scheduled initial phone interviews",
        "Administered pre-employment exams and coordinated medical exams",
        "Prepared job postings and updated the applicant database"],
    "Talent Acquisition Associate": [
        "Recruited for customer service and technical support accounts; hired 30+ agents per month",
        "Conducted phone screening, versant-style English assessments and endorsement to hiring managers",
        "Organized weekend walk-in hiring events"],
    "Customer Service Representative": [
        "Handled 80-100 inbound calls daily for a telecom and e-commerce account",
        "Consistently met AHT, CSAT and quality targets; Top Agent award (Q2)",
        "Resolved billing concerns and escalated complex issues to Tier 2"],
    "Technical Support Representative": [
        "Provided Tier 1 technical support for home internet and mobile devices",
        "Troubleshot connectivity issues and logged tickets in the CRM",
        "Maintained 92% first-call resolution rate"],
    "Team Leader, Customer Care": [
        "Leads a team of 18 customer service agents on a US healthcare account",
        "Conducts daily huddles, coaching sessions and performance reviews",
        "Monitors real-time adherence and coordinates with Workforce and Quality teams"],
    "Accounting Assistant": [
        "Prepared vouchers, check disbursements and bank reconciliations",
        "Recorded daily sales and collections in the accounting system",
        "Assisted in the preparation of monthly VAT and expanded withholding tax returns"],
    "Accounts Payable Clerk": [
        "Processed supplier invoices and prepared payment schedules",
        "Matched purchase orders, receiving reports and invoices",
        "Maintained the AP aging report and supplier records"],
    "Bookkeeper": [
        "Maintains books of accounts for 12 small business clients",
        "Prepares financial statements, VAT returns and percentage tax returns",
        "Reconciles bank statements and records petty cash replenishments"],
    "Executive Assistant": [
        "Manages calendar, travel and meetings of the Chief Operating Officer",
        "Prepares board meeting materials, minutes and correspondence",
        "Coordinates with department heads on deliverables and deadlines"],
    "Administrative Assistant": [
        "Handled office supplies inventory, purchase requests and vendor coordination",
        "Received and routed correspondence; maintained filing systems",
        "Coordinated meeting room bookings and company events"],
    "Training Specialist": [
        "Facilitates new hire training and product refresher courses for 100+ agents monthly",
        "Develops training modules, assessments and certification exams",
        "Tracks nesting performance and graduation rates of training batches"],
    "Office Manager": [
        "Oversees office administration, facilities and building maintenance",
        "Manages the office budget, supplier contracts and company assets",
        "Supervises 3 admin staff and the reception desk"],
    "Receptionist": [
        "Greeted visitors, managed the front desk and answered the trunkline",
        "Logged incoming deliveries and scheduled conference rooms"],
    "Data Encoder": [
        "Encoded customer applications and transaction forms with 99.5% accuracy",
        "Verified encoded data against source documents and flagged discrepancies"],
    "HR Intern": ["Assisted in filing 201 documents and preparing onboarding kits",
                  "Helped organize the company job fair"],
    "Office Intern": ["Assisted in encoding and filing documents", "Answered phone calls and prepared memos"],
}

ROLE_CATEGORY_SKILLS = {
    "payroll": ["Payroll processing (semi-monthly / monthly)", "SSS, PhilHealth & Pag-IBIG remittances",
                "TRAIN law withholding tax computation", "BIR Forms 1601-C, 1604-C, 2316", "MS Excel (VLOOKUP, pivot tables, macros)",
                "HRIS payroll modules", "Timekeeping and DTR validation", "Final pay & 13th month pay computation"],
    "hrops": ["201 file management", "HRIS administration", "Onboarding & offboarding", "DOLE labor standards",
              "MS Excel (VLOOKUP, pivot tables)", "Government compliance (SSS, PhilHealth, Pag-IBIG)", "Report writing"],
    "recruit": ["Full-cycle recruitment", "Sourcing via job portals and social media", "Behavioral interviewing",
                "Applicant tracking systems", "Employer branding", "MS Office"],
    "cs": ["Customer handling", "Active listening", "CRM tools", "Typing speed: 45 WPM", "Problem solving",
           "Fluent in English and Filipino"],
    "acct": ["Bookkeeping", "Bank reconciliation", "Accounts payable", "VAT and EWT returns", "QuickBooks / Xero",
             "MS Excel"],
    "admin": ["Calendar management", "Office administration", "Records management", "MS Office / Google Workspace",
              "Business correspondence", "Event coordination"],
}

TRAININGS = {
    "payroll": ["Seminar on TRAIN Law Withholding Tax Computation", "Payroll Masterclass: Philippine Statutory Deductions",
                "SSS, PhilHealth and Pag-IBIG Employer Compliance Briefing"],
    "hrops": ["Labor Standards and DOLE Compliance Workshop", "HR Analytics Using Excel"],
    "recruit": ["Behavioral Event Interviewing Workshop", "Employer Branding for Recruiters"],
    "cs": ["Customer Service Excellence Workshop", "Basic Occupational Safety and Health Orientation"],
    "acct": ["BIR eFPS and eBIRForms Orientation", "Basic Accounting Refresher Course"],
    "admin": ["Business Writing Workshop", "Records Management Seminar"],
}

DEGREES = {
    "payroll": ["BS Accountancy", "BS Business Administration major in Financial Management",
                "BS Accounting Technology"],
    "hrops": ["BS Psychology", "BS Business Administration major in Human Resource Development Management"],
    "recruit": ["BS Psychology", "AB Communication"],
    "cs": ["BS Hospitality Management", "AB English", "BS Information Technology"],
    "acct": ["BS Accountancy", "BS Accounting Technology"],
    "admin": ["BS Office Administration", "BS Business Administration major in Marketing Management"],
}


def yrange(start, end):
    """start/end: (year, month or None). end None = Present."""
    def f(t):
        y, m = t
        return f"{MONTHS[m - 1][:3]} {y}" if m else f"{y}"
    s = f(start)
    e = "Present" if end is None else f(end)
    if start[1] is None and (end is None or end[1] is None):
        return f"{s}{EN}{e}"
    return f"{s} {EN} {e}"


# Applicant plan. Exactly three (Reyes, Santos, Cruz) have 5+ years in payroll.
APPLICANT_PLAN = [
    # demo answers
    dict(key="reyes", tmpl=3, cat="payroll", skills_cat="hrops", ids=True, sex="F", first="Kristine Joy", middle="Manalo", last="Reyes",
         target="Payroll Supervisor", jobs=[
             ("HR Operations Supervisor", (2024, None), None),
             ("Payroll Specialist", (2017, None), (2024, None), "PAGE_BREAK"),
             ("HR Assistant", (2015, None), (2017, None))],
         objective="To obtain a supervisory role in a growing organization where I can apply my experience in HR "
                   "operations, compensation and government compliance, and contribute to accurate and timely "
                   "employee services."),
    dict(key="santos", tmpl=1, cat="payroll", ids=False, sex="M", first="Rodel", middle="Aguilar", last="Santos",
         target="Payroll Officer", jobs=[
             ("HR Generalist", (2022, 1), None),
             ("Payroll Officer", (2015, 6), (2021, 12))],
         summary="HR professional with more than six years of hands-on payroll experience (Payroll Officer, "
                 "Jun 2015 – Dec 2021) plus four years as an HR Generalist. Strong in statutory deductions, "
                 "TRAIN-law withholding tax and government remittances. Looking to return to a dedicated payroll role."),
    dict(key="cruz", tmpl=2, cat="payroll", ids=True, sex="F", first="Patricia Anne", middle="Domingo", last="Cruz",
         target="Payroll Lead", jobs=[
             ("Senior Payroll Analyst", (2019, None), (2025, None)),
             ("Payroll Associate", (2016, None), (2019, None))],
         summary="Payroll analyst with 9 years of Philippine payroll experience (Payroll Associate 2016–2019, "
                 "Senior Payroll Analyst 2019–2025) supporting up to 2,300 employees. Known for clean audits, "
                 "Excel automation and calm handling of payroll cut-offs."),
    # distractors: some payroll exposure, but < 5 years
    dict(tmpl=2, cat="admin", ids=True, target="Payroll Assistant", jobs=[
        ("Payroll Assistant", (2022, None), (2024, None)),
        ("Administrative Assistant", (2019, None), (2022, None))],
         summary="Detail-oriented administrative professional with two years of payroll support experience "
                 "(Payroll Assistant, 2022–2024). Eager to grow into a full payroll role."),
    dict(tmpl=3, cat="admin", ids=True, target="Timekeeping / Payroll Staff", jobs=[
        ("Timekeeping and Payroll Clerk", (2023, 3), None),
        ("Data Encoder", (2021, 6), (2023, 2))]),
    # others: no payroll
    dict(tmpl=1, cat="recruit", ids=False, target="Senior Recruitment Specialist", jobs=[
        ("Recruitment Specialist", (2020, None), None), ("Recruitment Assistant", (2018, None), (2020, None))]),
    dict(tmpl=2, cat="recruit", ids=False, target="Talent Acquisition Specialist", jobs=[
        ("Talent Acquisition Associate", (2023, None), None),
        ("Customer Service Representative", (2021, None), (2023, None))]),
    dict(tmpl=3, cat="hrops", ids=True, target="HR Assistant", jobs=[
        ("HR Assistant", (2022, None), None), ("Receptionist", (2020, None), (2022, None))]),
    dict(tmpl=1, cat="cs", ids=False, target="Customer Service Team Leader", jobs=[
        ("Customer Service Representative", (2019, None), None)]),
    dict(tmpl=2, cat="cs", ids=True, target="Technical Support Specialist", jobs=[
        ("Technical Support Representative", (2021, None), (2024, None)),
        ("Customer Service Representative", (2019, None), (2021, None))]),
    dict(tmpl=3, cat="cs", ids=True, target="Operations Supervisor", jobs=[
        ("Team Leader, Customer Care", (2020, None), None),
        ("Customer Service Representative", (2016, None), (2020, None))]),
    dict(tmpl=1, cat="acct", ids=True, target="Accounting Staff", jobs=[
        ("Accounting Assistant", (2020, None), None), ("Accounts Payable Clerk", (2018, None), (2020, None))]),
    dict(tmpl=2, cat="acct", ids=False, target="Bookkeeper / Accounting Associate", jobs=[
        ("Bookkeeper", (2017, None), None)]),
    dict(tmpl=3, cat="admin", ids=False, target="Executive Assistant", jobs=[
        ("Executive Assistant", (2019, None), None), ("Administrative Assistant", (2016, None), (2019, None))]),
    dict(tmpl=1, cat="cs", ids=False, target="Training Specialist", jobs=[
        ("Training Specialist", (2021, None), None),
        ("Customer Service Representative", (2018, None), (2021, None))]),
    dict(tmpl=2, cat="admin", ids=False, target="Office Manager", jobs=[
        ("Office Manager", (2018, None), None), ("Receptionist", (2015, None), (2018, None))]),
    dict(tmpl=3, cat="admin", ids=True, target="Administrative Assistant", jobs=[
        ("Administrative Assistant", (2023, None), None), ("Office Intern", (2022, 6), (2022, 9))]),
    dict(tmpl=1, cat="acct", ids=False, target="Accounts Payable Associate", jobs=[
        ("Accounts Payable Clerk", (2022, None), None)]),
    dict(tmpl=2, cat="admin", ids=False, target="Data Entry / Admin Staff", jobs=[
        ("Data Encoder", (2020, None), (2025, None))]),
    dict(tmpl=3, cat="hrops", ids=True, target="HR Generalist", jobs=[
        ("HR Assistant", (2021, None), None), ("HR Intern", (2020, 6), (2020, 10))]),
]

AREA_WORDS = {"payroll": "payroll", "hrops": "human resources operations", "recruit": "recruitment",
              "cs": "customer service", "acct": "accounting and bookkeeping", "admin": "office administration"}


def build_applicant(plan, idx):
    p = make_person(R, sex=plan.get("sex"), first=plan.get("first"), middle=plan.get("middle"),
                    last=plan.get("last"), yob=(1986, 2001))
    if plan.get("last") in ("Reyes", "Santos", "Cruz"):
        p["dob"] = {"Reyes": date(1991, 4, 18), "Santos": date(1988, 11, 2), "Cruz": date(1990, 7, 25)}[plan["last"]]
    p.update(plan)
    jobs = []
    used = set()
    for j in plan["jobs"]:
        title, start, end = j[:3]
        pool = [c for c in PRIOR_COMPANIES if c not in used]
        comp = R.choice(pool)
        used.add(comp)
        city = R.choice(LOCALES)[0]
        bl = ROLE_BULLETS[title][:]
        nb = len(bl) if title in ("Payroll Specialist", "Senior Payroll Analyst", "Payroll Officer",
                                  "HR Operations Supervisor") else min(len(bl), R.randint(2, 4))
        jobs.append(dict(title=title, company=comp, city=city, dates=yrange(start, end),
                         bullets=bl[:nb], page_break=len(j) > 3))
    p["job_list"] = jobs
    cat = plan["cat"]
    first_start = min(j[1][0] for j in plan["jobs"])
    grad = first_start - (0 if R.random() < 0.5 else 1)
    p["edu"] = [(R.choice(DEGREES[cat]), R.choice(SCHOOLS), f"{grad - 4}{EN}{grad}")]
    p["hs"] = (R.choice(HIGH_SCHOOLS), f"{grad - 8}{EN}{grad - 4}")
    if p["dob"].year > grad - 20:
        p["dob"] = p["dob"].replace(year=grad - 21, day=min(p["dob"].day, 28))
    scat = plan.get("skills_cat", cat)
    p["skills"] = ROLE_CATEGORY_SKILLS[scat][:R.randint(5, len(ROLE_CATEGORY_SKILLS[scat]))]
    p["trainings"] = [(t, R.randint(max(first_start, 2016), 2025)) for t in TRAININGS[cat]]
    if "summary" not in plan:
        yrs = TODAY.year - first_start
        adj = R.choice(["Dedicated", "Results-oriented", "Reliable", "Hardworking", "Detail-oriented"])
        noun = {"payroll": "payroll professional", "hrops": "HR practitioner", "recruit": "recruiter",
                "cs": "customer service professional", "acct": "accounting professional",
                "admin": "administrative professional"}[cat]
        p["summary"] = (f"{adj} {noun} with {yrs} year{'s' if yrs != 1 else ''} of experience in "
                        f"{AREA_WORDS[cat]}. Seeking a {plan['target']} position where I can contribute to "
                        f"team goals and continue to grow professionally.")
    if "objective" not in plan:
        p["objective"] = p["summary"]
    p["expected_salary"] = R.choice([None, None, 22000, 25000, 28500, 32000, 35000, 40000])
    if p.get("key"):
        p["expected_salary"] = {"reyes": 45000, "santos": 38000, "cruz": 52000}[p["key"]]
    p["height"] = f"{R.randint(150, 178)} cm"
    p["weight"] = f"{R.randint(48, 82)} kg"
    p["religion"] = R.choice(["Roman Catholic", "Roman Catholic", "Christian", "Iglesia ni Cristo", "Born Again Christian"])
    refs = []
    for _ in range(2):
        rp = make_person(R)
        refs.append(dict(name=f"{'Ms.' if rp['sex'] == 'F' else 'Mr.'} {rp['first']} {rp['last']}",
                         raw=f"{rp['first']} {rp['last']}", pos=R.choice(["HR Manager", "Operations Manager",
                                                                          "Finance Supervisor", "Team Manager",
                                                                          "Accounting Manager"]),
                         company=R.choice(PRIOR_COMPANIES), phone=gen_phone(False, R)))
        _used_surnames.discard(rp["last"])  # references may share surnames with others; fine
    p["refs"] = refs
    return p


def resume_personal_rows(d: Doc, p, with_ids=None):
    rows = [("Date of Birth", d.mark("DOB", fmt_long(p["dob"]))),
            ("Age", str(age_on(p["dob"]))),
            ("Civil Status", p["civil"]),
            ("Citizenship", "Filipino")]
    if with_ids if with_ids is not None else p["ids"]:
        rows += [("SSS No.", d.mark("SSS", p["sss"])), ("TIN", d.mark("TIN", p["tin"])),
                 ("PhilHealth No.", d.mark("PHILHEALTH", p["philhealth"])),
                 ("Pag-IBIG MID No.", d.mark("PAGIBIG", p["pagibig"]))]
    return rows


def resume_t1(p) -> Doc:
    """Classic serif, centered header."""
    d = Doc(f"Resume - {p['first']} {p['last']}", ml=60, mr=60, mt=56,
            on_new_page=lambda dd: (dd.text(dd.W - dd.mr, 36, dd.mark("NAME", f"{p['first']} {p['last']}") + " — Resume", "serif-i", 8.5, GRAY, "r"),))
    cx = d.W / 2
    d.text(cx, d.y + 22, d.mark("NAME", f"{p['first']} {p['middle'][0]}. {p['last']}".upper()), "serif-b", 22, BLACK, "c")
    d.y += 34
    d.text(cx, d.y + 10, d.mark("ADDRESS", p["addr"]), "serif", 9.5, DARK, "c")
    d.y += 14
    d.text(cx, d.y + 10, f"{d.mark('PHONE', p['phone'])}   •   {d.mark('EMAIL', p['email'])}", "serif", 9.5, DARK, "c")
    d.y += 18
    d.line(d.ml, d.y, d.W - d.mr, d.y, 1.0, DARK)
    d.y += 10

    def section(t):
        d.ensure(40)
        d.y += 6
        d.text(d.ml, d.y + 11, t.upper(), "serif-b", 11.5, BLACK)
        d.y += 15
        d.line(d.ml, d.y, d.W - d.mr, d.y, 0.5, RULE)
        d.y += 7

    section("Professional Summary")
    d.para(p["summary"], "serif", 10.5)
    section("Work Experience")
    for j in p["job_list"]:
        d.ensure(60)
        d.text(d.ml, d.y + 11, j["title"], "serif-b", 11)
        d.y += 15
        d.text(d.ml, d.y + 10, f"{j['company']}, {j['city']}  |  {j['dates']}", "serif-i", 10, DARK)
        d.y += 15
        for b in j["bullets"]:
            d.bullet(b, "serif", 10.2)
        d.y += 5
    section("Education")
    for deg, school, yrs in p["edu"]:
        d.text(d.ml, d.y + 11, deg, "serif-b", 10.5)
        d.text(d.W - d.mr, d.y + 11, yrs, "serif", 10.5, DARK, "r")
        d.y += 14
        d.text(d.ml, d.y + 10, school, "serif-i", 10, DARK)
        d.y += 16
    section("Skills")
    d.para(" • ".join(p["skills"]), "serif", 10.2)
    section("Personal Details")
    for k, v in resume_personal_rows(d, p):
        d.ensure(15)
        d.text(d.ml, d.y + 10, k, "serif-b", 10)
        d.text(d.ml + 120, d.y + 10, v, "serif", 10)
        d.y += 14.5
    if p["expected_salary"]:
        d.text(d.ml, d.y + 10, "Expected Salary", "serif-b", 10)
        d.text(d.ml + 120, d.y + 10, d.mark("SALARY", peso(p["expected_salary"])) + " / month", "serif", 10)
        d.y += 14.5
    section("References")
    d.para("Available upon request.", "serif", 10.5)
    return d


def resume_t2(p) -> Doc:
    """Modern teal header band, sans."""
    def cont(dd):
        dd.rect(0, 0, dd.W, 8, fill=TEAL)
        dd.text(dd.ml, 30, f"{p['first']} {p['last']}", "sans-b", 9, TEAL)
        dd.y = 48
    d = Doc(f"CV - {p['first']} {p['last']}", ml=50, mr=50, mt=40, on_new_page=cont)
    d.rect(0, 0, d.W, 112, fill=TEAL)
    d.text(d.ml, 46, d.mark("NAME", f"{p['first']} {p['last']}"), "sans-b", 24, WHITE)
    d.text(d.ml, 66, p["target"], "sans", 12, (0.85, 0.95, 0.95))
    d.text(d.ml, 86, f"{d.mark('PHONE', p['phone'])}   |   {d.mark('EMAIL', p['email'])}", "sans", 9.5, WHITE)
    d.text(d.ml, 101, d.mark("ADDRESS", p["addr"]), "sans", 9.5, WHITE)
    d.y = 128

    def section(t):
        d.ensure(40)
        d.y += 6
        d.text(d.ml, d.y + 11, t.upper(), "sans-b", 10.5, TEAL)
        d.y += 16
        d.line(d.ml, d.y, d.ml + 40, d.y, 2, TEAL)
        d.y += 8

    section("Profile")
    d.para(p["summary"], "sans", 9.8)
    section("Experience")
    for j in p["job_list"]:
        d.ensure(60)
        d.text(d.ml, d.y + 10.5, j["title"], "sans-b", 10.5)
        d.y += 14
        d.text(d.ml, d.y + 9.5, f"{j['company']}  ·  {j['city']}  ·  {j['dates']}", "sans", 9.3, GRAY)
        d.y += 14
        for b in j["bullets"]:
            d.bullet(b, "sans", 9.6, sym="–")
        d.y += 6
    section("Education")
    for deg, school, yrs in p["edu"]:
        d.text(d.ml, d.y + 10, deg, "sans-b", 10)
        d.y += 13.5
        d.text(d.ml, d.y + 9.5, f"{school}  ·  {yrs}", "sans", 9.5, GRAY)
        d.y += 16
    section("Skills")
    sk = p["skills"]
    half = (len(sk) + 1) // 2
    y0 = d.y
    for col, items in enumerate((sk[:half], sk[half:])):
        d.y = y0
        for s in items:
            d.bullet(s, "sans", 9.5, x=d.ml + col * (d.cw / 2))
    d.y = y0 + half * 13.2 + 4
    section("Personal Information")
    rows = resume_personal_rows(d, p)
    if p["expected_salary"]:
        rows.append(("Expected Salary", d.mark("SALARY", peso(p["expected_salary"]))))
    colw = d.cw / 2
    for i in range(0, len(rows), 2):
        d.ensure(15)
        for c, (k, v) in enumerate(rows[i:i + 2]):
            d.text(d.ml + c * colw, d.y + 9.5, f"{k}:", "sans-b", 9.3, DARK)
            d.text(d.ml + c * colw + 98, d.y + 9.5, v, "sans", 9.3)
        d.y += 14
    return d


def resume_t3(p) -> Doc:
    """Filipino biodata-style resume (Verdana, photo box, personal information block)."""
    d = Doc(f"Resume of {p['first']} {p['last']}", ml=54, mr=54, mt=50,
            on_new_page=lambda dd: (dd.text(dd.ml, 34, "Resume – " + dd.mark("NAME", f"{p['first']} {p['middle'][0]}. {p['last']}"), "verd", 7.5, GRAY),
                                    dd.line(dd.ml, 40, dd.W - dd.mr, 40, 0.5, RULE)))
    d.rect(d.W - d.mr - 72, d.y, 72, 72, stroke=GRAY)
    d.text(d.W - d.mr - 36, d.y + 34, "2x2", "verd", 8, LIGHT, "c")
    d.text(d.W - d.mr - 36, d.y + 45, "PHOTO", "verd", 8, LIGHT, "c")
    d.text(d.ml, d.y + 16, d.mark("NAME", f"{p['first']} {p['middle']} {p['last']}".upper()), "verd-b", 15)
    d.y += 30
    for ln in wrap(d.mark("ADDRESS", p["addr"]), "verd", 8.6, d.cw - 90):
        d.text(d.ml, d.y + 9, ln, "verd", 8.6, DARK)
        d.y += 12.5
    d.text(d.ml, d.y + 9, f"Mobile No.: {d.mark('PHONE', p['phone'])}", "verd", 8.6, DARK)
    d.y += 12.5
    d.text(d.ml, d.y + 9, f"E-mail: {d.mark('EMAIL', p['email'])}", "verd", 8.6, DARK)
    d.y = max(d.y + 10, 50 + 82)

    def section(t):
        d.ensure(42)
        d.y += 6
        d.rect(d.ml, d.y, d.cw, 16, fill=BAND)
        d.text(d.ml + 6, d.y + 11.5, t.upper(), "verd-b", 9.2, NAVY)
        d.y += 23

    def kv(k, v, lw=130):
        d.ensure(14)
        d.text(d.ml + 6, d.y + 9, k, "verd", 8.8, DARK)
        d.text(d.ml + lw, d.y + 9, ":", "verd", 8.8, DARK)
        lines = wrap(v, "verd", 8.8, d.cw - lw - 12)
        for i, ln in enumerate(lines):
            if i:
                d.y += 12.6
            d.text(d.ml + lw + 10, d.y + 9, ln, "verd", 8.8)
        d.y += 13.2

    section("Objective")
    d.para(p["objective"], "verd", 8.9, x=d.ml + 6)
    section("Personal Information")
    kv("Date of Birth", d.mark("DOB", fmt_long(p["dob"])))
    kv("Age", str(age_on(p["dob"])))
    kv("Place of Birth", p["pob"])
    kv("Sex", "Female" if p["sex"] == "F" else "Male")
    kv("Civil Status", p["civil"])
    kv("Citizenship", "Filipino")
    kv("Religion", p["religion"])
    kv("Height / Weight", f"{p['height']} / {p['weight']}")
    if p["ids"]:
        kv("SSS No.", d.mark("SSS", p["sss"]))
        kv("TIN", d.mark("TIN", p["tin"]))
        kv("PhilHealth No.", d.mark("PHILHEALTH", p["philhealth"]))
        kv("Pag-IBIG MID No.", d.mark("PAGIBIG", p["pagibig"]))
    if p["expected_salary"]:
        kv("Expected Salary", d.mark("SALARY", peso(p["expected_salary"])))
    section("Educational Attainment")
    deg, school, yrs = p["edu"][0]
    kv("Tertiary", f"{deg}, {school} ({yrs})")
    kv("Secondary", f"{p['hs'][0]} ({p['hs'][1]})")
    section("Skills")
    for s in p["skills"]:
        d.bullet(s, "verd", 8.8, x=d.ml + 6)
    section("Work Experience")
    for j in p["job_list"]:
        if j["page_break"] and len(d.pages) == 1:
            d.new_page()
        d.ensure(60)
        d.text(d.ml + 6, d.y + 10, f"{j['title']}, {j['dates']}", "verd-b", 9.2)
        d.y += 14
        d.text(d.ml + 6, d.y + 9, f"{j['company']} – {j['city']}", "verd-i", 8.6, DARK)
        d.y += 14
        for b in j["bullets"]:
            d.bullet(b, "verd", 8.8, x=d.ml + 10)
        d.y += 6
    section("Trainings / Seminars Attended")
    for t, y in p["trainings"]:
        d.bullet(f"{t} ({y})", "verd", 8.8, x=d.ml + 6)
    section("Character References")
    for r in p["refs"]:
        d.ensure(30)
        d.text(d.ml + 6, d.y + 9, d.mark("NAME", r["name"]), "verd-b", 8.8)
        d.y += 12.5
        d.text(d.ml + 6, d.y + 9, f"{r['pos']}, {r['company']}  |  {d.mark('PHONE', r['phone'])}", "verd", 8.6, DARK)
        d.y += 16
    d.ensure(70)
    d.y += 8
    d.para("I hereby certify that the above information is true and correct to the best of my knowledge and belief.",
           "verd-i", 8.6, x=d.ml + 6, color=DARK)
    d.y += 26
    nm = f"{p['first']} {p['middle'][0]}. {p['last']}".upper()
    d.line(d.W - d.mr - 190, d.y, d.W - d.mr, d.y, 0.6, DARK)
    d.text(d.W - d.mr - 95, d.y + 12, d.mark("NAME", nm), "verd-b", 8.6, BLACK, "c")
    d.text(d.W - d.mr - 95, d.y + 24, "Applicant", "verd", 7.8, GRAY, "c")
    d.y += 30
    return d


# ---------------------------------------------------------------------------
# 201 files, payslips
# ---------------------------------------------------------------------------
DEPTS = [("Operations", "Customer Service Representative"), ("Operations", "Customer Service Representative"),
         ("Operations", "Senior Customer Service Representative"), ("Operations", "Team Leader"),
         ("Operations", "Operations Manager"), ("Human Resources", "HR Generalist"),
         ("Human Resources", "Recruitment Specialist"), ("Finance", "Accounting Associate"),
         ("Information Technology", "IT Support Specialist"), ("Quality Assurance", "Quality Analyst"),
         ("Workforce Management", "Real-Time Analyst"), ("Training", "Trainer"),
         ("Administration", "Administrative Assistant"), ("Operations", "Technical Support Representative"),
         ("Finance", "Billing Specialist")]
SAL_BY_POS = {"Customer Service Representative": (18000, 24000), "Senior Customer Service Representative": (24000, 29000),
              "Team Leader": (32000, 40000), "Operations Manager": (60000, 75000), "HR Generalist": (28000, 34000),
              "Recruitment Specialist": (24000, 30000), "Accounting Associate": (23000, 28000),
              "IT Support Specialist": (26000, 33000), "Quality Analyst": (25000, 30000),
              "Real-Time Analyst": (25000, 30000), "Trainer": (30000, 38000), "Administrative Assistant": (18000, 22000),
              "Technical Support Representative": (21000, 26000), "Billing Specialist": (22000, 27000)}


def build_employee(i):
    p = make_person(R, yob=(1980, 2002))
    dept, pos = DEPTS[i]
    lo, hi = SAL_BY_POS[pos]
    p["salary"] = R.randrange(lo, hi + 1, 500)
    hired = rand_date(2015, 2025) if i != 10 else date(2026, 10, 1)
    p["hired"] = hired
    p["empno"] = f"BOC-{hired.year}-{R.randint(100, 999):04d}"
    p["dept"], p["pos"] = dept, pos
    p["status"] = "Probationary" if (TODAY - hired).days < 183 else "Regular"
    p["bank"] = gen_bank_account()
    if p["civil"] == "Married":
        sp_first = R.choice(MALE if p["sex"] == "F" else FEMALE)
        p["ec_name"], p["ec_rel"] = f"{sp_first} {p['last']}", "Spouse"
        kids = []
        for _ in range(R.randint(1, 2)):
            ks = R.choice("FM")
            kd = rand_date(max(p["dob"].year + 22, 2008), 2023)
            kids.append((f"{R.choice(FEMALE if ks == 'F' else MALE)} {p['last']}", kd, "Child"))
        p["dependents"] = [(f"{sp_first} {p['last']}", rand_date(p['dob'].year - 3, p['dob'].year + 3), "Spouse")] + kids
    else:
        p["ec_name"], p["ec_rel"] = f"{R.choice(FEMALE)} {p['last']}", "Mother"
        p["dependents"] = []
    p["ec_phone"] = gen_phone(R.random() < 0.3)
    if R.random() < 0.4:
        p["perm_addr"] = gen_address(R)[0]
    else:
        p["perm_addr"] = None
    p["blood"] = R.choice(["O+", "A+", "B+", "AB+", "O+"])
    return p


def form_row(d: Doc, cells, h=27):
    """cells: [(label, value, frac)] -> boxed form row."""
    d.ensure(h)
    x = d.ml
    for label, value, frac in cells:
        w = d.cw * frac
        d.rect(x, d.y, w, h, stroke=(0.55, 0.55, 0.58), lw=0.5)
        d.text(x + 4, d.y + 8.5, label, "sans", 6.3, GRAY)
        size = 9.6
        while size > 6.5 and sw(value, "sans", size) > w - 8:
            size -= 0.3
        d.text(x + 4, d.y + 21, value, "sans", size)
        x += w
    d.y += h


def form_band(d: Doc, title, color=NAVY):
    d.ensure(48)
    d.y += 8
    d.rect(d.ml, d.y, d.cw, 15, fill=color)
    d.text(d.ml + 5, d.y + 10.8, title.upper(), "sans-b", 8.5, WHITE)
    d.y += 15


def company_header(d: Doc, right_title, right_sub=None):
    d.rect(d.ml, d.y, 30, 30, fill=TEAL)
    d.text(d.ml + 15, d.y + 20, "BOC", "sans-b", 9, WHITE, "c")
    d.text(d.ml + 38, d.y + 13, COMPANY, "sans-b", 13, NAVY)
    d.text(d.ml + 38, d.y + 26, COMPANY_ADDR, "sans", 7.4, GRAY)
    d.y += 44
    d.text(d.W / 2, d.y + 14, right_title, "sans-b", 14, BLACK, "c")
    d.y += 20
    if right_sub:
        d.text(d.W / 2, d.y + 9, right_sub, "sans", 8.5, GRAY, "c")
        d.y += 14


def file_201(p) -> Doc:
    d = Doc(f"201 File - {p['last']}, {p['first']}", ml=48, mr=48, mt=44, mb=60)
    company_header(d, "EMPLOYEE PERSONAL DATA SHEET (201 FILE)", "CONFIDENTIAL — For HR use only")
    form_band(d, "I. Employee Information")
    form_row(d, [("EMPLOYEE NO.", p["empno"], 0.25), ("LAST NAME", d.mark("NAME", p["last"].upper()), 0.25),
                 ("FIRST NAME", d.mark("NAME", p["first"].upper()), 0.27),
                 ("MIDDLE NAME", d.mark("NAME", p["middle"].upper()), 0.23)])
    form_row(d, [("DATE OF BIRTH (MM/DD/YYYY)", d.mark("DOB", fmt_mdy(p["dob"])), 0.25),
                 ("PLACE OF BIRTH", p["pob"], 0.27), ("SEX", "Female" if p["sex"] == "F" else "Male", 0.14),
                 ("CIVIL STATUS", p["civil"], 0.17), ("BLOOD TYPE", p["blood"], 0.17)])
    form_row(d, [("PRESENT ADDRESS", d.mark("ADDRESS", p["addr"]), 1.0)])
    form_row(d, [("PERMANENT / PROVINCIAL ADDRESS",
                  d.mark("ADDRESS", p["perm_addr"]) if p["perm_addr"] else "Same as present address", 1.0)])
    form_row(d, [("MOBILE NO.", d.mark("PHONE", p["phone"]), 0.3), ("PERSONAL E-MAIL", d.mark("EMAIL", p["email"]), 0.45),
                 ("CITIZENSHIP", "Filipino", 0.25)])
    form_band(d, "II. Government Numbers")
    form_row(d, [("SSS NO.", d.mark("SSS", p["sss"]), 0.25), ("TIN", d.mark("TIN", p["tin"]), 0.25),
                 ("PHILHEALTH NO.", d.mark("PHILHEALTH", p["philhealth"]), 0.25),
                 ("PAG-IBIG MID NO.", d.mark("PAGIBIG", p["pagibig"]), 0.25)])
    form_band(d, "III. Employment Details")
    form_row(d, [("DATE HIRED", fmt_long(p["hired"]), 0.25), ("EMPLOYMENT STATUS", p["status"], 0.25),
                 ("POSITION / DESIGNATION", p["pos"], 0.5)])
    form_row(d, [("DEPARTMENT", p["dept"], 0.35), ("BASIC MONTHLY SALARY", d.mark("SALARY", peso(p["salary"])), 0.3),
                 ("PAY FREQUENCY", "Semi-monthly", 0.35)])
    form_row(d, [("PAYROLL BANK", PAYROLL_BANK, 0.35), ("PAYROLL ACCOUNT NO.", d.mark("BANK_ACCOUNT", p["bank"]), 0.3),
                 ("ACCOUNT TYPE", "Savings (payroll)", 0.35)])
    form_band(d, "IV. Family / Dependents")
    if p["dependents"]:
        for nm, dob, rel in p["dependents"]:
            form_row(d, [("NAME", d.mark("NAME", nm), 0.5), ("RELATIONSHIP", rel, 0.25),
                         ("DATE OF BIRTH", d.mark("DOB", fmt_mdy(dob)), 0.25)])
    else:
        form_row(d, [("NAME", "None declared", 0.5), ("RELATIONSHIP", "—", 0.25), ("DATE OF BIRTH", "—", 0.25)])
    form_band(d, "V. Person to Contact in Case of Emergency")
    form_row(d, [("NAME", d.mark("NAME", p["ec_name"]), 0.45), ("RELATIONSHIP", p["ec_rel"], 0.2),
                 ("CONTACT NO.", d.mark("PHONE", p["ec_phone"]), 0.35)])
    d.y += 14
    d.para("I certify that the information above is true and complete. I authorize Bayanihan Outsourcing Corp. to "
           "process my personal data for employment, payroll and statutory compliance purposes in accordance with "
           "the Data Privacy Act of 2012.", "sans", 7.8, color=DARK)
    d.y += 30
    sig = f"{p['first']} {p['middle'][0]}. {p['last']}"
    d.mark("NAME", sig)
    d.line(d.ml, d.y, d.ml + 200, d.y, 0.6, DARK)
    d.text(d.ml + 100, d.y + 12, sig, "sans-b", 9, BLACK, "c")
    d.text(d.ml + 100, d.y + 22, "Employee Signature over Printed Name", "sans", 7, GRAY, "c")
    d.line(d.W - d.mr - 150, d.y, d.W - d.mr, d.y, 0.6, DARK)
    d.text(d.W - d.mr - 75, d.y + 12, fmt_long(p["hired"] - timedelta(days=3)), "sans", 9, BLACK, "c")
    d.text(d.W - d.mr - 75, d.y + 22, "Date Accomplished", "sans", 7, GRAY, "c")
    return d


def compute_payslip(p, rng):
    monthly = p["salary"]
    semi = monthly / 2
    hourly = monthly * 12 / 261 / 8
    ot_hrs = rng.choice([0, 0, 2, 4, 6, 8, 10.5, 12])
    ot = round(hourly * 1.25 * ot_hrs, 2)
    nd_hrs = rng.choice([0, 0, 16, 24, 40]) if p["dept"] == "Operations" else 0
    nd = round(hourly * 0.10 * nd_hrs, 2)
    rice = 1000.00
    msc = min(max(round(monthly / 500) * 500, 5000), 35000)
    sss = round(msc * 0.05 / 2, 2)
    ph = round(min(max(monthly, 10000), 100000) * 0.025 / 2, 2)
    hdmf = 100.00
    taxable = semi + ot + nd - sss - ph - hdmf
    if taxable <= 10417:
        wtax = 0.0
    elif taxable <= 16666:
        wtax = (taxable - 10417) * 0.15
    elif taxable <= 33332:
        wtax = 937.50 + (taxable - 16667) * 0.20
    elif taxable <= 83332:
        wtax = 4270.70 + (taxable - 33333) * 0.25
    else:
        wtax = 16770.70 + (taxable - 83333) * 0.30
    wtax = round(wtax, 2)
    loan = rng.choice([0, 0, 0, 1083.33, 875.00])
    earnings = [(f"Basic Pay (semi-monthly)", round(semi, 2)), (f"Overtime ({ot_hrs:g} hrs @ 125%)", ot),
                (f"Night Differential ({nd_hrs} hrs)", nd), ("Rice Subsidy (de minimis)", rice)]
    deductions = [("SSS Contribution", sss), ("PhilHealth Contribution", ph), ("Pag-IBIG Contribution", hdmf),
                  ("Withholding Tax", wtax)]
    if loan:
        deductions.append(("SSS Salary Loan", loan))
    gross = round(sum(a for _, a in earnings), 2)
    tded = round(sum(a for _, a in deductions), 2)
    return earnings, deductions, gross, tded, round(gross - tded, 2)


def payslip(p, rng) -> Doc:
    d = Doc(f"Payslip - {p['empno']} - Sep 16-30 2026", ml=48, mr=48, mt=44, mb=60)
    company_header(d, "PAYSLIP", f"Pay Period: September 16{EN}30, 2026     Pay Date: September 30, 2026")
    full = f"{p['last'].upper()}, {p['first']} {p['middle'][0]}."
    form_row(d, [("EMPLOYEE NAME", d.mark("NAME", full), 0.5), ("EMPLOYEE NO.", p["empno"], 0.25),
                 ("TAX STATUS", "S" if p["civil"] == "Single" else "ME", 0.25)])
    form_row(d, [("POSITION", p["pos"], 0.5), ("DEPARTMENT", p["dept"], 0.5)])
    form_row(d, [("TIN", d.mark("TIN", p["tin"]), 0.25), ("SSS NO.", d.mark("SSS", p["sss"]), 0.25),
                 ("PHILHEALTH NO.", d.mark("PHILHEALTH", p["philhealth"]), 0.25),
                 ("PAG-IBIG MID NO.", d.mark("PAGIBIG", p["pagibig"]), 0.25)])
    form_row(d, [("MONTHLY BASIC RATE", d.mark("SALARY", peso(p["salary"])), 0.33),
                 ("CREDITED TO", f"{PAYROLL_BANK} Payroll", 0.33),
                 ("ACCOUNT NO.", d.mark("BANK_ACCOUNT", p["bank"]), 0.34)])
    earnings, deductions, gross, tded, net = compute_payslip(p, rng)
    d.y += 12
    half = d.cw / 2 - 6
    xs = (d.ml, d.ml + half + 12)
    y0 = d.y
    for col, (title, items) in enumerate((("EARNINGS", earnings), ("DEDUCTIONS", deductions))):
        x = xs[col]
        y = y0
        d.rect(x, y, half, 16, fill=BAND)
        d.text(x + 5, y + 11.3, title, "sans-b", 8.5, NAVY)
        d.text(x + half - 5, y + 11.3, "AMOUNT", "sans-b", 8.5, NAVY, "r")
        y += 16
        for name, amt in items:
            d.text(x + 5, y + 13, name, "sans", 9)
            amt_s = peso(amt)
            if name.startswith("Basic Pay"):
                d.mark("SALARY", amt_s)
            d.text(x + half - 5, y + 13, amt_s, "sans", 9, BLACK, "r")
            d.line(x, y + 18, x + half, y + 18, 0.4, RULE)
            y += 18
        total_lbl = "GROSS PAY" if col == 0 else "TOTAL DEDUCTIONS"
        total = gross if col == 0 else tded
        y = y0 + 16 + 18 * 5 + 4
        d.text(x + 5, y + 13, total_lbl, "sans-b", 9.3)
        ts = peso(total)
        if col == 0:
            d.mark("SALARY", ts)
        d.text(x + half - 5, y + 13, ts, "sans-b", 9.3, BLACK, "r")
        d.line(x, y + 1, x + half, y + 1, 0.8, DARK)
    d.y = y0 + 16 + 18 * 5 + 30
    d.rect(d.ml, d.y, d.cw, 30, fill=(0.93, 0.97, 0.96), stroke=TEAL, lw=0.8)
    d.text(d.ml + 10, d.y + 19.5, "NET PAY", "sans-b", 12, TEAL)
    d.text(d.W - d.mr - 10, d.y + 19.5, d.mark("SALARY", peso(net)), "sans-b", 13, TEAL, "r")
    d.y += 44
    d.para("This is a system-generated payslip and does not require a signature. Please report any discrepancy to "
           "the Payroll Unit within five (5) working days from pay date. Keep this payslip confidential.",
           "sans-i", 7.8, color=GRAY)
    return d


# ---------------------------------------------------------------------------
# DOCX
# ---------------------------------------------------------------------------
def new_docx(title):
    doc = docx.Document()
    cp = doc.core_properties
    cp.title, cp.author, cp.comments = title, "Gel demo data generator", FOOTER
    from datetime import datetime
    cp.created = cp.modified = datetime(2026, 9, 25, 9, 0, 0)
    st = doc.styles["Normal"]
    st.font.name, st.font.size = "Arial", Pt(10.5)
    fp = doc.sections[0].footer.paragraphs[0]
    run = fp.add_run(FOOTER)
    run.font.size, run.font.color.rgb = Pt(7), RGBColor(0x99, 0x99, 0x9C)
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    hp = doc.sections[0].header.paragraphs[0]
    hr = hp.add_run(f"{COMPANY}  |  {COMPANY_ADDR}")
    hr.font.size, hr.font.color.rgb = Pt(7.5), RGBColor(0x0D, 0x5C, 0x66)
    return doc


def docx_text(doc) -> str:
    parts = [p.text for p in doc.paragraphs]
    for t in doc.tables:
        for row in t.rows:
            for c in row.cells:
                parts.append(c.text)
    return re.sub(r"\s+", " ", " ".join(parts))


def emit_docx(doc, pii, path: Path):
    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(path))
    record(str(path.relative_to(OUT)), pii, [docx_text(doc)])


def heading(doc, text, size=14, center=True):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.bold, r.font.size = True, Pt(size)
    if center:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    return p


# ---------------------------------------------------------------------------
# Personal pack
# ---------------------------------------------------------------------------
def personal_person():
    p = make_person(R, sex="F", first="Camille", middle="Esguerra", last="Villafuerte", yob=(1993, 1993))
    p["dob"] = date(1993, 2, 14)
    p["phone"] = gen_phone(False)
    p["email"] = "camille.villafuerte@example-mail.ph"
    p["psn"] = gen_psn()
    p["passport"] = gen_passport()
    p["dl"] = gen_dl()
    p["acct"] = gen_bank_account()
    p["card"] = gen_card("4000")
    p["cc"] = gen_card("5200")
    p["gcash"] = gen_phone(False)
    p["employer"] = "Anahaw Digital Ventures Inc."
    p["salary"] = 58500
    return p


def kv_sheet(d: Doc, rows, lw=170, size=10):
    for k, v in rows:
        d.ensure(18)
        d.text(d.ml, d.y + 11, k, "sans", size - 0.6, GRAY)
        lines = wrap(v, "sans", size, d.cw - lw)
        for i, ln in enumerate(lines):
            if i:
                d.y += 13
            d.text(d.ml + lw, d.y + 11, ln, "sans", size)
        d.y += 13
        d.line(d.ml, d.y + 2, d.W - d.mr, d.y + 2, 0.4, RULE)
        d.y += 6


def doc_header(d: Doc, org, sub, title, color=NAVY):
    d.text(d.ml, d.y + 14, org, "sans-b", 14, color)
    d.text(d.ml, d.y + 27, sub, "sans", 8, GRAY)
    d.y += 36
    d.line(d.ml, d.y, d.W - d.mr, d.y, 1.2, color)
    d.y += 12
    d.text(d.ml, d.y + 14, title, "sans-b", 13)
    d.y += 26


def bank_statement(p) -> Doc:
    d = Doc("Bangko Halimbawa - Statement of Account", ml=48, mr=48)
    doc_header(d, "BANGKO HALIMBAWA", "A fictional bank for demonstration purposes  ·  Customer care: (02) 8888 0000",
               "STATEMENT OF ACCOUNT — Savings", TEAL)
    kv_sheet(d, [("Account Name", d.mark("NAME", p["full"].upper())),
                 ("Mailing Address", d.mark("ADDRESS", p["addr"])),
                 ("Account Number", d.mark("BANK_ACCOUNT", p["acct"])),
                 ("Linked Debit Card No.", d.mark("CARD", p["card"])),
                 ("Statement Period", f"September 1{EN}30, 2026"),
                 ("Branch", "Halimbawa Cubao Branch")], lw=150)
    d.y += 8
    d.rect(d.ml, d.y, d.cw, 16, fill=BAND)
    for x, t, a in [(d.ml + 4, "DATE", "l"), (d.ml + 60, "DESCRIPTION", "l"), (d.ml + 360, "DEBIT", "r"),
                    (d.ml + 430, "CREDIT", "r"), (d.W - d.mr - 4, "BALANCE", "r")]:
        d.text(x, d.y + 11, t, "sans-b", 8, NAVY, a)
    d.y += 20
    rng = random.Random(SEED + 77)
    bal = 41235.60
    sal_half = p["salary"] / 2
    tx = [(1, "Balance brought forward", None, None)]
    tx += [(2, "ATM Withdrawal – Cubao", 3000, None),
           (3, f"InstaPay Transfer to GCash {p['gcash']}", 2500, None),
           (5, "Bills Payment – Liwanag Power Co.", 3184.25, None),
           (7, "POS Purchase – Palengke Mart Demo", 1872.40, None),
           (9, "Bills Payment – Agos Water Services", 642.10, None),
           (12, "Fund Transfer to Acct 418-2290-117 (Rent)", 12000, None),
           (14, f"Payroll Credit – {p['employer']}", None, sal_half),
           (15, f"Credit Card Payment – Card {p['cc']}", 8500, None),
           (17, "POS Purchase – Tindahan Online Demo", 2399, None),
           (18, f"InstaPay Transfer to GCash {p['gcash']}", 1500, None),
           (20, "ATM Withdrawal – Ortigas", 5000, None),
           (22, "Interest Credit", None, 3.42),
           (22, "Withholding Tax on Interest", 0.68, None),
           (24, "Bills Payment – Kidlat Fiber Internet", 1699, None),
           (26, "POS Purchase – Botika Halimbawa", 845.75, None),
           (27, "Transfer from Acct 230-5518-0042", None, 1500),
           (28, "Online Purchase – Lakbay Halimbawa Travel", 6250, None),
           (30, f"Payroll Credit – {p['employer']}", None, sal_half)]
    for day, desc, deb, cred in tx:
        if deb:
            bal -= deb
        if cred:
            bal += cred
        d.ensure(16)
        d.text(d.ml + 4, d.y + 9, f"09/{day:02d}/2026", "sans", 8.4)
        d.text(d.ml + 60, d.y + 9, desc, "sans", 8.4)
        if deb:
            d.text(d.ml + 360, d.y + 9, f"{deb:,.2f}", "sans", 8.4, BLACK, "r")
        if cred:
            s = f"{cred:,.2f}"
            d.text(d.ml + 430, d.y + 9, s, "sans", 8.4, BLACK, "r")
        d.text(d.W - d.mr - 4, d.y + 9, f"{bal:,.2f}", "sans", 8.4, BLACK, "r")
        d.line(d.ml, d.y + 13, d.W - d.mr, d.y + 13, 0.3, RULE)
        d.y += 16
    d.mark("SALARY", f"{sal_half:,.2f}")
    d.mark("GCASH", p["gcash"])
    d.mark("CARD", p["cc"])
    d.y += 12
    kv_sheet(d, [("Ending Balance", peso(bal)), ("Total Payroll Credits", d.mark("SALARY", peso(sal_half * 2)))], lw=150)
    d.y += 10
    d.para("Please examine this statement upon receipt. If no error is reported within thirty (30) days, the account "
           "will be considered correct. Deposits are insured up to the maximum amount allowed by law. Bangko "
           "Halimbawa is a fictional institution created for software demonstration.", "sans-i", 7.6, color=GRAY)
    return d


def passport_details(p) -> Doc:
    d = Doc("Visa Assistance - Applicant Passport Details", ml=54, mr=54)
    issue = date(2022, 6, 20)
    doc_header(d, "LAKBAY HALIMBAWA TRAVEL & TOURS", "Visa assistance desk (fictional agency)  ·  Form VA-02",
               "Visa Application — Applicant Passport Details", NAVY)
    kv_sheet(d, [("Surname", d.mark("NAME", p["last"].upper())), ("Given Names", p["first"].upper()),
                 ("Middle Name", p["middle"].upper()),
                 ("Date of Birth", d.mark("DOB", fmt_long(p["dob"]))), ("Place of Birth", p["pob"]),
                 ("Sex", "F"), ("Nationality", "Filipino"),
                 ("Passport Number", d.mark("PASSPORT", p["passport"])),
                 ("Date of Issue", fmt_long(issue)), ("Date of Expiry", fmt_long(issue.replace(year=issue.year + 10) - timedelta(days=1))),
                 ("Place of Issue", "Manila"),
                 ("Home Address", d.mark("ADDRESS", p["addr"])),
                 ("Mobile Number", d.mark("PHONE", p["phone"])), ("E-mail Address", d.mark("EMAIL", p["email"])),
                 ("Occupation / Employer", f"Senior Product Analyst, {p['employer']}"),
                 ("Monthly Income", d.mark("SALARY", peso(p["salary"]))),
                 ("Destination / Travel Dates", "Japan (Osaka), December 18–26, 2026"),
                 ("Emergency Contact", f"{d.mark('NAME', 'Divina Villafuerte')} (Mother) – {d.mark('PHONE', '+63 917 482 0391')}")])
    d.y += 14
    d.para("Applicant's declaration: I certify that the passport details above are true and match my valid passport.",
           "sans-i", 8.5, color=DARK)
    d.y += 26
    d.line(d.ml, d.y, d.ml + 210, d.y, 0.6, DARK)
    d.text(d.ml + 105, d.y + 12, d.mark("NAME", p["full"]), "sans-b", 9, BLACK, "c")
    return d


def drivers_license_details(p) -> Doc:
    d = Doc("Motor Insurance Application - Driver Details", ml=54, mr=54)
    doc_header(d, "KALASAG HALIMBAWA INSURANCE", "Motor car insurance (fictional insurer)  ·  Application No. MC-2026-08817",
               "Principal Driver — Driver's License Details", (0.45, 0.12, 0.12))
    kv_sheet(d, [("Full Name", d.mark("NAME", p["lfm"])),
                 ("Date of Birth", d.mark("DOB", fmt_mdy(p["dob"]))),
                 ("Address", d.mark("ADDRESS", p["addr"])),
                 ("Driver's License No.", d.mark("DRIVERS_LICENSE", p["dl"])),
                 ("License Type", "Non-Professional"), ("DL Codes / Restrictions", "A, B, B1"),
                 ("License Expiry", "February 14, 2031"), ("Conditions", "None"),
                 ("Mobile Number", d.mark("PHONE", p["phone"])), ("E-mail", d.mark("EMAIL", p["email"])),
                 ("Vehicle", "2021 Toyota-class compact sedan (example), Plate No. NDX 4521"),
                 ("Years of Driving Experience", "8"),
                 ("Claims in Last 3 Years", "None")])
    d.y += 14
    d.para("The applicant declares that the license details above are accurate and that the license is valid and not "
           "suspended.", "sans-i", 8.5, color=DARK)
    return d


def philsys_details(p) -> Doc:
    d = Doc("Bangko Halimbawa - Customer Information Sheet", ml=54, mr=54)
    doc_header(d, "BANGKO HALIMBAWA", "Customer Information Sheet (KYC update)  ·  fictional bank",
               "Customer Information Sheet — Individual", TEAL)
    kv_sheet(d, [("Customer Name", d.mark("NAME", p["full"])),
                 ("Date of Birth", d.mark("DOB", fmt_long(p["dob"]))),
                 ("Place of Birth", p["pob"]), ("Civil Status", p["civil"]), ("Nationality", "Filipino"),
                 ("Primary ID Presented", "Philippine Identification (PhilSys)"),
                 ("PhilSys Number (PSN)", d.mark("PHILSYS", p["psn"])),
                 ("TIN", d.mark("TIN", p["tin"])),
                 ("SSS No.", d.mark("SSS", p["sss"])),
                 ("Present Address", d.mark("ADDRESS", p["addr"])),
                 ("Mobile Number", d.mark("PHONE", p["phone"])),
                 ("E-wallet (GCash) Number", d.mark("GCASH", p["gcash"])),
                 ("E-mail", d.mark("EMAIL", p["email"])),
                 ("Employer", p["employer"]), ("Position", "Senior Product Analyst"),
                 ("Gross Monthly Income", d.mark("SALARY", peso(p["salary"]))),
                 ("Source of Funds", "Salary"),
                 ("Existing Account No.", d.mark("BANK_ACCOUNT", p["acct"]))])
    d.y += 14
    d.para("I certify that the information provided is true and correct and I consent to its processing for "
           "know-your-customer purposes.", "sans-i", 8.5, color=DARK)
    return d


def medical_certificate(p) -> Doc:
    d = Doc("Medical Certificate", ml=64, mr=64)
    d.text(d.W / 2, d.y + 16, "KLINIKA HALIMBAWA MEDICAL CENTER", "serif-b", 16, (0.1, 0.3, 0.5), "c")
    d.text(d.W / 2, d.y + 30, "45 Mabini St., Brgy. Kamuning, Quezon City  ·  Tel. (02) 8700 1234  ·  fictional clinic",
           "serif", 8.5, GRAY, "c")
    d.y += 42
    d.line(d.ml, d.y, d.W - d.mr, d.y, 1.0, (0.1, 0.3, 0.5))
    d.y += 28
    d.text(d.W / 2, d.y + 16, "MEDICAL CERTIFICATE", "serif-b", 17, BLACK, "c")
    d.y += 40
    d.text(d.W - d.mr, d.y + 11, "Date: September 22, 2026", "serif", 11, BLACK, "r")
    d.y += 30
    d.para("TO WHOM IT MAY CONCERN:", "serif-b", 11.5)
    d.y += 8
    age = age_on(p["dob"], date(2026, 9, 22))
    d.mark("NAME", p["full"].upper())
    d.mark("ADDRESS", p["addr"])
    d.mark("DOB", fmt_long(p["dob"]))
    d.mark("PHILHEALTH", p["philhealth"])
    d.para(f"This is to certify that {p['full'].upper()}, {age} years old, female, born on {fmt_long(p['dob'])}, of "
           f"{p['addr']}, was seen and examined at this clinic on September 22, 2026.",
           "serif", 11.5, leading=19)
    d.y += 8
    kv_sheet(d, [("PhilHealth No.", p["philhealth"]), ("Diagnosis", "Acute gastroenteritis, mild dehydration"),
                 ("Recommendation", "Rest for three (3) days, September 22–24, 2026; oral rehydration; follow-up as needed"),
                 ("Fit to work on", "September 25, 2026")], lw=130, size=10.5)
    d.y += 10
    d.para("This certification is issued upon the request of the patient for whatever legal purpose it may serve, "
           "except medico-legal.", "serif", 11.5, leading=19)
    d.y += 50
    doc_name = "Dr. Reynaldo M. Abad, MD"
    d.mark("NAME", doc_name)
    d.line(d.W - d.mr - 200, d.y, d.W - d.mr, d.y, 0.6, DARK)
    d.text(d.W - d.mr - 100, d.y + 14, doc_name, "serif-b", 11, BLACK, "c")
    d.text(d.W - d.mr - 100, d.y + 27, "Attending Physician  ·  PRC Lic. No. 0098765", "serif", 9, GRAY, "c")
    return d


# ---------------------------------------------------------------------------
# Phone-scan simulation
# ---------------------------------------------------------------------------
def _solve(A, b):
    n = len(b)
    M = [row[:] + [b[i]] for i, row in enumerate(A)]
    for c in range(n):
        piv = max(range(c, n), key=lambda r: abs(M[r][c]))
        M[c], M[piv] = M[piv], M[c]
        for r in range(n):
            if r != c:
                f = M[r][c] / M[c][c]
                for k in range(c, n + 1):
                    M[r][k] -= f * M[c][k]
    return [M[i][n] / M[i][i] for i in range(n)]


def persp_coeffs(dst, src):
    A, B = [], []
    for (x, y), (X, Y) in zip(dst, src):
        A.append([x, y, 1, 0, 0, 0, -X * x, -X * y]); B.append(X)
        A.append([0, 0, 0, x, y, 1, -Y * x, -Y * y]); B.append(Y)
    return _solve(A, B)


def phone_scan(img: Image.Image, rng: random.Random) -> Image.Image:
    w, h = img.size
    # paper tint (off-white / slightly warm)
    tint = (rng.randint(242, 250), rng.randint(236, 244), rng.randint(220, 232))
    img = ImageChops.multiply(img, Image.new("RGB", img.size, tint))
    # uneven lighting: soft gradient, darker toward one side
    corners = [255, 255 - rng.randint(4, 14), 255 - rng.randint(8, 22), 255 - rng.randint(0, 10)]
    rng.shuffle(corners)
    g = Image.new("L", (2, 2))
    g.putdata(corners)
    g = g.resize((w, h), Image.Resampling.BILINEAR)
    img = ImageChops.multiply(img, Image.merge("RGB", (g, g, g)))
    # slight perspective, placed on a desk background
    pad = int(w * 0.035)
    W2, H2 = w + 2 * pad, h + 2 * pad
    j = lambda: rng.uniform(-0.012, 0.012) * w
    dst = [(pad + j(), pad + j()), (pad + w + j(), pad + j()), (pad + w + j(), pad + h + j()), (pad + j(), pad + h + j())]
    src = [(0, 0), (w, 0), (w, h), (0, h)]
    desk = rng.choice([(78, 70, 62), (96, 98, 104), (120, 104, 86)])
    out = img.transform((W2, H2), Image.Transform.PERSPECTIVE, persp_coeffs(dst, src), resample=Image.Resampling.BICUBIC, fillcolor=desk)
    out = out.rotate(rng.uniform(-1.5, 1.5), resample=Image.Resampling.BICUBIC, expand=False, fillcolor=desk)
    # crop as a scanner app would, leaving a thin uneven border
    cl = int(pad * rng.uniform(0.55, 0.85))
    out = out.crop((cl, int(pad * rng.uniform(0.55, 0.85)), W2 - int(pad * rng.uniform(0.55, 0.85)),
                    H2 - int(pad * rng.uniform(0.55, 0.85))))
    # mild sensor noise (deterministic)
    nw, nh = out.size
    noise = Image.frombytes("L", (nw, nh), rng.randbytes(nw * nh)).point(lambda v: 128 + (v - 128) * 0.10)
    noise = noise.filter(ImageFilter.GaussianBlur(0.6))
    out = ImageChops.add(out, Image.merge("RGB", (noise, noise, noise)), scale=1.0, offset=-128)
    out = out.filter(ImageFilter.GaussianBlur(rng.uniform(0.45, 0.7)))
    return out


def save_jpeg(img, path: Path, quality=75):
    path.parent.mkdir(parents=True, exist_ok=True)
    img.save(str(path), "JPEG", quality=quality, optimize=False)


def jpeg_bytes(img, quality=75):
    b = io.BytesIO()
    img.save(b, "JPEG", quality=quality)
    return b.getvalue()


def image_only_pdf(jpegs: list[bytes], path: Path, title):
    path.parent.mkdir(parents=True, exist_ok=True)
    c = rl_canvas.Canvas(str(path), pagesize=A4, invariant=1)
    c.setTitle(title)
    for jb in jpegs:
        ir = ImageReader(io.BytesIO(jb))
        iw, ih = ir.getSize()
        pw = A4[0]
        ph = pw * ih / iw
        c.setPageSize((pw, ph))
        c.drawImage(ir, 0, 0, width=pw, height=ph)
        c.showPage()
    c.save()


def record_scan(file_rel, doc: Doc, page_indices):
    texts = doc.page_texts()
    sub_texts = [texts[i] for i in page_indices]
    record(file_rel, doc.pii_on_pages(page_indices), sub_texts)


def pii_on_pages(self, idxs):
    texts = self.page_texts()
    return [(t, v) for t, v in self.pii if any(v in texts[i] for i in idxs)]


Doc.pii_on_pages = pii_on_pages


# ---------------------------------------------------------------------------
# Main
# ---------------------------------------------------------------------------
REGEX = {
    "SSS": r"(?<![\d-])\d{2}-\d{7}-\d(?![\d-])",
    "TIN": r"(?<![\d-])\d{3}-\d{3}-\d{3}-(?:\d{5}|\d{3})(?![\d-])",
    "PHILHEALTH": r"(?<![\d-])\d{2}-\d{9}-\d(?![\d-])",
    "PAGIBIG": r"(?<![\d-])\d{4}-\d{4}-\d{4}(?![\d-])",
    "PHILSYS": r"(?<![\d-])\d{4}-\d{4}-\d{4}-\d{4}(?![\d-])",
    "PASSPORT": r"\b[A-Z]\d{7}[A-Z]\b",
    "DRIVERS_LICENSE": r"\b[A-Z]\d{2}-\d{2}-\d{6}\b",
    "BANK_ACCOUNT": r"(?<![\d-])\d{3}-\d{4}-\d{3,5}(?![\d-])",
    "CARD": r"(?<!\d)\d{4} \d{4} \d{4} \d{4}(?!\d)",
    "PHONE": r"(?<!\d)(?:\+63 ?9\d{2}|09\d{2})[ -]?\d{3}[ -]?\d{4}(?!\d)",
    "GCASH": r"(?<!\d)(?:\+63 ?9\d{2}|09\d{2})[ -]?\d{3}[ -]?\d{4}(?!\d)",
    "EMAIL": r"[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+\.[A-Za-z]{2,}",
    "SALARY": "₱?\\s?\\d{1,3}(?:,\\d{3})*\\.\\d{2}",
    "DOB": r"(?:(?:January|February|March|April|May|June|July|August|September|October|November|December) \d{1,2}, (?:19|20)\d{2}|\d{2}/\d{2}/(?:19|20)\d{2})",
    "ADDRESS": r".+, (?:Brgy\. ).+ \d{4}$",
    "NAME": r"^[A-Za-zÑñ .,'-]+$",
}


def main():
    if OUT.exists():
        shutil.rmtree(OUT)
    OUT.mkdir(parents=True)

    # ---------------- Resumes ----------------
    applicants = [build_applicant(pl, i) for i, pl in enumerate(APPLICANT_PLAN)]
    order = list(range(len(applicants)))
    R.shuffle(order)  # so demo answers are not files 1-3
    resume_docs = {}
    name_styles = [lambda p: f"{p['last']}_{p['first'].split()[0]}_Resume.pdf",
                   lambda p: f"CV - {p['first']} {p['last']}.pdf",
                   lambda p: f"{p['first'].replace(' ', '')}{p['last'].replace(' ', '')}_Resume_2026.pdf",
                   lambda p: f"Resume_{p['last'].upper().replace(' ', '_')}.pdf"]
    demo = []
    for n, i in enumerate(order):
        p = applicants[i]
        d = {1: resume_t1, 2: resume_t2, 3: resume_t3}[p["tmpl"]](p)
        fname = name_styles[(n + i) % 4](p)
        path = HR / "Resumes" / fname
        emit_pdf(d, path)
        resume_docs[i] = (d, path)
        if p.get("key"):
            pages = [k + 1 for k, t in enumerate(d.page_texts()) if "Payroll" in t and any(
                j["title"] in t for j in p["job_list"] if "Payroll" in j["title"])]
            pj = [f"{j['title']}, {j['company']}, {j['dates']}" for j in p["job_list"] if "Payroll" in j["title"]]
            demo.append({"surname": p["last"], "name": f"{p['first']} {p['middle']} {p['last']}",
                         "file": str(path.relative_to(OUT)), "payroll_experience": pj,
                         "payroll_experience_pages": pages, "total_pages": len(d.pages)})

    # ---------------- 201 files + payslips ----------------
    employees = [build_employee(i) for i in range(15)]
    f201 = {}
    for e in employees:
        d = file_201(e)
        path = HR / "201 Files" / f"201_{e['empno']}_{e['last'].replace(' ', '')}_{e['first'].split()[0]}.pdf"
        emit_pdf(d, path)
        f201[e["empno"]] = (d, path)
    pay = {}
    for k, e in enumerate(employees[:10]):
        d = payslip(e, random.Random(SEED + 100 + k))
        path = HR / "Payslips" / f"Payslip_2026-09-30_{e['empno']}_{e['last'].replace(' ', '')}.pdf"
        emit_pdf(d, path)
        pay[e["empno"]] = (d, path)

    # ---------------- Contracts (DOCX) ----------------
    hr_dir = make_person(R, sex="F", first="Liza")
    hr_dir_name = f"{hr_dir['first']} {hr_dir['middle'][0]}. {hr_dir['last']}"
    e = employees[10]
    doc = new_docx("Contract of Employment")
    pii = []
    heading(doc, "CONTRACT OF EMPLOYMENT (PROBATIONARY)", 15)
    doc.add_paragraph("KNOW ALL MEN BY THESE PRESENTS:")
    doc.add_paragraph(
        f"This Contract of Employment is entered into this 25th day of September 2026 in Pasig City, by and between:")
    doc.add_paragraph(
        f"{COMPANY}, a corporation duly organized under Philippine law, with principal office at {COMPANY_ADDR}, "
        f"represented herein by its HR Director, {hr_dir_name} (the “COMPANY”);")
    doc.add_paragraph("— and —").alignment = WD_ALIGN_PARAGRAPH.CENTER
    emp_full = f"{e['first']} {e['middle']} {e['last']}".upper()
    doc.add_paragraph(
        f"{emp_full}, of legal age, {e['civil'].lower()}, Filipino, with residence at {e['addr']}, with SSS No. "
        f"{e['sss']} and TIN {e['tin']} (the “EMPLOYEE”).")
    pii += [("NAME", hr_dir_name), ("NAME", emp_full), ("ADDRESS", e["addr"]), ("SSS", e["sss"]), ("TIN", e["tin"])]
    heading(doc, "TERMS AND CONDITIONS", 11.5, center=False)
    terms = [
        f"Position. The EMPLOYEE is engaged as {e['pos']} under the {e['dept']} Department, effective "
        f"{fmt_long(e['hired'])}.",
        "Probationary Period. The EMPLOYEE shall undergo a probationary period of six (6) months, during which "
        "performance shall be evaluated against the reasonable standards made known at the time of engagement.",
        f"Compensation. The EMPLOYEE shall receive a basic monthly salary of {peso(e['salary'])}, payable "
        f"semi-monthly every 15th and 30th of the month through payroll account no. {e['bank']} with {PAYROLL_BANK}.",
        "Statutory Benefits. The EMPLOYEE shall be enrolled with SSS, PhilHealth and Pag-IBIG and shall be entitled "
        "to 13th month pay and service incentive leave as provided by law.",
        "Work Schedule. Forty (40) hours per week on a shifting schedule, including night shift, with night "
        "differential pay as provided by law.",
        "Confidentiality and Data Privacy. The EMPLOYEE shall keep confidential all client, customer and employee "
        "information and comply with the Data Privacy Act of 2012."]
    pii += [("SALARY", peso(e["salary"])), ("BANK_ACCOUNT", e["bank"])]
    for k, t in enumerate(terms, 1):
        doc.add_paragraph(f"{k}. {t}")
    doc.add_paragraph("IN WITNESS WHEREOF, the parties have signed this Contract on the date and place first above written.")
    t = doc.add_table(rows=3, cols=2)
    t.cell(0, 0).text, t.cell(0, 1).text = "For the COMPANY:", "EMPLOYEE:"
    t.cell(1, 0).text, t.cell(1, 1).text = "\n\n______________________________", "\n\n______________________________"
    t.cell(2, 0).text = f"{hr_dir_name}\nHR Director"
    t.cell(2, 1).text = f"{e['first']} {e['middle'][0]}. {e['last']}\nMobile: {e['phone']}"
    pii += [("NAME", f"{e['first']} {e['middle'][0]}. {e['last']}"), ("PHONE", e["phone"])]
    wit = make_person(R)
    doc.add_paragraph(f"Signed in the presence of: {wit['first']} {wit['last']} (HR Generalist)")
    pii.append(("NAME", f"{wit['first']} {wit['last']}"))
    contract_path = HR / "Contracts" / f"Employment_Contract_{e['last'].replace(' ', '')}_{e['first'].split()[0]}.docx"
    emit_docx(doc, pii, contract_path)

    # Offer letter to Reyes
    rey = next(a for a in applicants if a.get("key") == "reyes")
    doc = new_docx("Job Offer")
    pii = []
    doc.add_paragraph("September 28, 2026")
    rey_name = f"{rey['first']} {rey['middle'][0]}. {rey['last']}"
    doc.add_paragraph(f"MS. {rey_name.upper()}\n{rey['addr']}\n{rey['email']}  |  {rey['phone']}")
    pii += [("NAME", f"MS. {rey_name.upper()}"), ("ADDRESS", rey["addr"]), ("EMAIL", rey["email"]),
            ("PHONE", rey["phone"])]
    heading(doc, "Subject: Offer of Employment — Payroll Supervisor", 11.5, center=False)
    doc.add_paragraph(f"Dear Ms. {rey['last']},")
    doc.add_paragraph(
        f"We are pleased to offer you the position of Payroll Supervisor in the Finance & Payroll Unit of {COMPANY}, "
        f"reporting to the Finance Manager. Your expected start date is November 3, 2026.")
    doc.add_paragraph("Compensation and benefits:")
    offer_sal = peso(45000)
    for b in [f"Basic monthly salary: {offer_sal}, paid semi-monthly",
              f"Monthly allowance: {peso(3000)} (communication and transportation)",
              "HMO coverage from day one, plus one (1) free dependent upon regularization",
              "15 days vacation leave and 15 days sick leave upon regularization",
              "13th month pay and performance bonus per company policy"]:
        doc.add_paragraph(b, style="List Bullet")
    pii.append(("SALARY", offer_sal))
    doc.add_paragraph(
        f"This offer is contingent on the submission of your pre-employment requirements, including your SSS No. "
        f"({rey['sss']}), TIN ({rey['tin']}), PhilHealth No. ({rey['philhealth']}) and Pag-IBIG MID No. "
        f"({rey['pagibig']}) as declared in your application, an NBI clearance and a medical certificate.")
    pii += [("SSS", rey["sss"]), ("TIN", rey["tin"]), ("PHILHEALTH", rey["philhealth"]), ("PAGIBIG", rey["pagibig"])]
    doc.add_paragraph("Kindly signify your acceptance by signing below and returning a copy on or before October 5, 2026.")
    doc.add_paragraph(f"Sincerely,\n\n\n{hr_dir_name}\nHR Director, {COMPANY}")
    pii.append(("NAME", hr_dir_name))
    doc.add_paragraph(f"CONFORME:\n\n\n______________________________\n{rey_name}    Date: ____________")
    pii.append(("NAME", rey_name))
    emit_docx(doc, pii, HR / "Contracts" / f"Offer_Letter_{rey['last']}_{rey['first'].split()[0]}.docx")

    # HR memo
    doc = new_docx("HR Memo 2026-014")
    pii = []
    heading(doc, "MEMORANDUM", 16)
    meta = doc.add_table(rows=4, cols=2)
    for r_, (k, v) in enumerate([("MEMO NO.", "HR-2026-014"), ("TO", "Concerned Employees (see list below)"),
                                 ("FROM", f"{hr_dir_name}, HR Director"), ("DATE", "September 29, 2026")]):
        meta.cell(r_, 0).text, meta.cell(r_, 1).text = k, v
    pii.append(("NAME", hr_dir_name))
    heading(doc, "SUBJECT: Update of Government Numbers and Payroll Bank Enrollment", 11.5, center=False)
    doc.add_paragraph(
        "Our records show discrepancies between the government numbers on file and those reported in the latest "
        "SSS and PhilHealth remittance reports. To avoid posting errors in your contributions and delays in payroll "
        "crediting, the following employees are requested to visit the HR office with their original IDs on or "
        "before October 10, 2026:")
    tbl = doc.add_table(rows=1, cols=5)
    tbl.style = "Table Grid"
    for c_, h_ in enumerate(["Employee No.", "Name", "SSS No. on file", "PhilHealth No. on file", "Concern"]):
        tbl.rows[0].cells[c_].text = h_
    concerns = ["SSS number mismatch", "Unposted PhilHealth premium (Aug 2026)", "Payroll account not yet enrolled",
                "TIN not yet submitted (new hire)"]
    for e_, cn in zip([employees[2], employees[5], employees[11], employees[10]], concerns):
        row = tbl.add_row().cells
        nm = f"{e_['last']}, {e_['first']} {e_['middle'][0]}."
        row[0].text, row[1].text, row[2].text, row[3].text, row[4].text = e_["empno"], nm, e_["sss"], e_["philhealth"], cn
        pii += [("NAME", nm), ("SSS", e_["sss"]), ("PHILHEALTH", e_["philhealth"])]
    doc.add_paragraph(
        "Please also be reminded that payroll accounts must be under your own name. Third-party or joint accounts "
        "will not be credited. For questions, please approach the HR Operations desk at the 8th floor.")
    doc.add_paragraph("For strict compliance.")
    emit_docx(doc, pii, HR / "Contracts" / "HR_Memo_2026-014_Government_Numbers.docx")

    # ---------------- Scans ----------------
    scans_dir = HR / "Scans"
    scan_specs = [  # (kind, key, output, filename)
        ("pay", employees[0]["empno"], "jpg", "IMG_20260930_191204.jpg"),
        ("pay", employees[3]["empno"], "jpg", "IMG_20260930_191322.jpg"),
        ("201", employees[6]["empno"], "jpg", "IMG_20260912_093415.jpg"),
        ("201", employees[8]["empno"], "jpg", "IMG_20260912_093502.jpg"),
        ("res", [i for i in order if applicants[i]["tmpl"] == 2 and not applicants[i].get("key")][1], "jpg",
         "IMG_20260918_141027.jpg"),
        ("res", [i for i in order if applicants[i]["tmpl"] == 1 and not applicants[i].get("key")][0], "pdf",
         "Scan 2026-09-18 14.12.pdf"),
        ("res", [i for i in order if applicants[i]["tmpl"] == 3 and applicants[i]["target"].startswith("Timekeeping")][0],
         "pdf", "Scan 2026-09-18 14.15.pdf"),
        ("201", employees[12]["empno"], "pdf", "Scan 2026-09-12 09.40.pdf"),
        ("201", employees[14]["empno"], "pdf", "Scan 2026-09-12 09.43.pdf"),
        ("pay", employees[7]["empno"], "pdf", "Scan 2026-10-01 08.05.pdf"),
    ]
    scan_sources = []
    for n, (kind, key, fmt, fname) in enumerate(scan_specs):
        d, src = {"pay": pay, "201": f201, "res": resume_docs}[kind][key]
        rng = random.Random(SEED + 1000 + n)
        idxs = [0] if fmt == "jpg" else list(range(len(d.pages)))
        imgs = [phone_scan(render_page_image(d.pages[i], dpi=rng.choice([160, 170, 180])), rng) for i in idxs]
        path = scans_dir / fname
        if fmt == "jpg":
            save_jpeg(imgs[0], path)
        else:
            image_only_pdf([jpeg_bytes(im) for im in imgs], path, "Scanned document")
        record_scan(str(path.relative_to(OUT)), d, idxs)
        scan_sources.append({"scan": str(path.relative_to(OUT)), "source": str(src.relative_to(OUT)),
                             "pages": len(idxs)})

    # ---------------- Personal pack ----------------
    pp = personal_person()
    pdir = OUT / "Personal"
    pdocs = {}
    for fn, builder in [("bank_statement.pdf", bank_statement), ("passport_details.pdf", passport_details),
                        ("drivers_license_details.pdf", drivers_license_details),
                        ("philsys_details.pdf", philsys_details), ("medical_certificate.pdf", medical_certificate)]:
        d = builder(pp)
        emit_pdf(d, pdir / fn)
        pdocs[fn] = d
    for n, (src, out) in enumerate([("medical_certificate.pdf", "medical_certificate_scan.jpg"),
                                    ("philsys_details.pdf", "philsys_details_scan.jpg")]):
        rng = random.Random(SEED + 2000 + n)
        d = pdocs[src]
        save_jpeg(phone_scan(render_page_image(d.pages[0], dpi=170), rng), pdir / out)
        record_scan(f"Personal/{out}", d, [0])

    # ---------------- Clipboard samples ----------------
    cdir = OUT / "clipboard-samples"
    cdir.mkdir(parents=True)
    e = employees[4]
    clip = {}
    clip["employee_record.txt"] = (
        f"Employee: {e['first']} {e['middle']} {e['last']} ({e['empno']})\n"
        f"Position: {e['pos']} - {e['dept']}\n"
        f"DOB: {fmt_long(e['dob'])}\n"
        f"Address: {e['addr']}\n"
        f"SSS: {e['sss']} | TIN: {e['tin']} | PhilHealth: {e['philhealth']} | Pag-IBIG: {e['pagibig']}\n"
        f"Basic monthly salary: {peso(e['salary'])}\n"
        f"Payroll acct ({PAYROLL_BANK}): {e['bank']}\n"
        f"Mobile: {e['phone']}  Email: {e['email']}\n",
        [("NAME", f"{e['first']} {e['middle']} {e['last']}"), ("DOB", fmt_long(e["dob"])), ("ADDRESS", e["addr"]),
         ("SSS", e["sss"]), ("TIN", e["tin"]), ("PHILHEALTH", e["philhealth"]), ("PAGIBIG", e["pagibig"]),
         ("SALARY", peso(e["salary"])), ("BANK_ACCOUNT", e["bank"]), ("PHONE", e["phone"]), ("EMAIL", e["email"])])
    a = next(x for x in applicants if x.get("key") == "santos")
    clip["resume_snippet.txt"] = (
        f"{a['first']} {a['middle'][0]}. {a['last']}\n{a['addr']}\n{a['phone']} | {a['email']}\n\n"
        f"Payroll Officer, {a['job_list'][1]['company']} ({a['job_list'][1]['dates']})\n"
        f"- Owned end-to-end payroll processing for 650 employees\n"
        f"Date of Birth: {fmt_long(a['dob'])}   Civil Status: {a['civil']}\n",
        [("NAME", f"{a['first']} {a['middle'][0]}. {a['last']}"), ("ADDRESS", a["addr"]), ("PHONE", a["phone"]),
         ("EMAIL", a["email"]), ("DOB", fmt_long(a["dob"]))])
    clip["passport_bank_snippet.txt"] = (
        f"Hi! Here are my details for the booking:\n"
        f"Name: {pp['full']}\nPassport No.: {pp['passport']} (exp. June 19, 2032)\n"
        f"Birthday: {fmt_long(pp['dob'])}\n\n"
        f"For the downpayment you can send to:\n{PAYROLL_BANK} Savings Acct No. {pp['acct']}\n"
        f"or GCash {pp['gcash']} ({pp['first']} {pp['last'][0]}.)\n"
        f"Card on file for the balance: {pp['cc']}\n",
        [("NAME", pp["full"]), ("PASSPORT", pp["passport"]), ("DOB", fmt_long(pp["dob"])),
         ("BANK_ACCOUNT", pp["acct"]), ("GCASH", pp["gcash"]), ("CARD", pp["cc"])])
    clip["clean_paragraph.txt"] = (
        "Reminder for the Q4 team planning session: we will review the onboarding checklist, agree on the new "
        "ticket triage flow, and finalize the training calendar for November. Please bring your top three process "
        "improvement ideas. The session runs from 2:00 PM to 4:30 PM in Conference Room B, and snacks will be "
        "provided. No preparation is needed beyond reading the one-page agenda shared last week.\n", [])
    for fn, (txt, pii) in clip.items():
        (cdir / fn).write_text(txt, encoding="utf-8")
        record(f"clipboard-samples/{fn}", pii, [re.sub(r"\s+", " ", txt)])

    # ---------------- Ground truth + README ----------------
    gt = {
        "description": "Ground truth for Gel's synthetic demo data. ALL values are fictional.",
        "seed": SEED,
        "page_note": "page is 1-based. DOCX and .txt files have no fixed pagination, so page is always 1 for them. "
                     "JPG scans are single images (page 1). A value printed on several pages is listed once per page.",
        "salary_note": "SALARY covers compensation figures (monthly/basic rate, expected salary, gross, net pay, "
                       "payroll credits) — not individual deduction line items.",
        "regex": REGEX,
        "demo_answers": {
            "query": "Sino sa applicants ang may 5+ years sa payroll?",
            "applicants": sorted(demo, key=lambda x: ["Reyes", "Santos", "Cruz"].index(x["surname"])),
            "near_misses": [f"{a['first']} {a['last']} ({', '.join(j['title'] + ' ' + j['dates'] for j in a['job_list'] if 'Payroll' in j['title'])})"
                            for a in applicants if not a.get("key") and any("Payroll" in j["title"] for j in a["job_list"])],
        },
        "scans": scan_sources,
        "entities": ENTITIES,
    }
    (OUT / "ground_truth.json").write_text(json.dumps(gt, indent=2, ensure_ascii=False), encoding="utf-8")
    write_readme()
    verify(gt)


def write_readme():
    (OUT / "README.md").write_text(f"""# Gel demo data (synthetic)

**All data in this folder is synthetic.** Every person, address, phone number, e-mail, government ID,
bank account, card number, company and bank is fictional and generated by
`scripts/generate_demo_data.py` (seed `{SEED}`). Each page carries the footer
"{FOOTER}". Fictional organisations used: {COMPANY}, Bangko Halimbawa,
Lakbay Halimbawa Travel & Tours, Kalasag Halimbawa Insurance, Klinika Halimbawa Medical Center.
ID numbers are format-valid but random; card numbers are Luhn-valid test-style numbers (4000… / 5200…).
Phone numbers use real Philippine mobile prefixes with random digits, so a collision with a real
number is possible; never call or message them.

## Layout

| Path | Contents |
|---|---|
| `HR Files/Resumes/` | 20 applicant resumes (PDF with text layer, 3 layouts) |
| `HR Files/201 Files/` | 15 employee personal data sheets (PDF) |
| `HR Files/Payslips/` | 10 payslips, Sep 16–30 2026 cut-off (PDF) |
| `HR Files/Contracts/` | employment contract, offer letter, HR memo (DOCX) |
| `HR Files/Scans/` | 10 phone-scanned copies: 5 JPG + 5 image-only PDF (no text layer, OCR required) |
| `Personal/` | bank statement, passport / driver's license / PhilSys details, medical certificate (PDF) + 2 scanned JPGs |
| `clipboard-samples/` | text snippets for Leak Guard testing (3 with PII, 1 clean) |
| `ground_truth.json` | every planted PII value: `entities[] = {{file, type, value, page}}`, plus `demo_answers` and the `regex` per type |

## Demo query

"Sino sa applicants ang may 5+ years sa payroll?" → exactly three applicants: **Reyes, Santos, Cruz**
(see `demo_answers` in `ground_truth.json`; the Reyes payroll experience is on page 2 of her resume).
Two near-miss applicants have under 5 years of payroll experience.

## Regenerate

```sh
python3 -m venv scripts/.venv
scripts/.venv/bin/pip install -r scripts/requirements.txt
scripts/.venv/bin/python scripts/generate_demo_data.py
```

Needs macOS (system Arial / Times New Roman / Verdana / Courier New fonts for the ₱ sign).
Output is deterministic. Optional text-layer / OCR check: `swift scripts/extract_text.swift pdf <file.pdf>`
or `swift scripts/extract_text.swift ocr <image.jpg>`.
""", encoding="utf-8")


def verify(gt):
    print("== File counts ==")
    for sub in ["HR Files/Resumes", "HR Files/201 Files", "HR Files/Payslips", "HR Files/Contracts", "HR Files/Scans",
                "Personal", "clipboard-samples"]:
        files = sorted(f for f in (OUT / sub).iterdir() if f.is_file())
        exts = {}
        for f in files:
            exts[f.suffix] = exts.get(f.suffix, 0) + 1
        print(f"  {sub}: {len(files)} {exts}")
    print("== Regex / Luhn validation ==")
    bad = 0
    counts = {}
    for ent in gt["entities"]:
        counts[ent["type"]] = counts.get(ent["type"], 0) + 1
        if not re.fullmatch(REGEX[ent["type"]], ent["value"]):
            print("  REGEX FAIL", ent)
            bad += 1
        if ent["type"] == "CARD" and not luhn_ok(ent["value"]):
            print("  LUHN FAIL", ent)
            bad += 1
    print(f"  {len(gt['entities'])} entities, {bad} failures; by type: {dict(sorted(counts.items()))}")
    print("== Image-only PDFs have no text operators / fonts ==")
    for f in sorted((OUT / "HR Files/Scans").glob("*.pdf")):
        raw = f.read_bytes()
        has_text = False
        # inspect every non-image stream (page content streams) for text-showing operators
        for m in re.finditer(rb"<<(.*?)>>\s*stream\r?\n(.*?)endstream", raw, re.S):
            if b"/Image" in m.group(1):
                continue
            body = m.group(2)
            if b"FlateDecode" in m.group(1):
                try:
                    body = zlib.decompress(body)
                except zlib.error:
                    pass
            if re.search(rb"\bBT\b|\bTj\b|\bTJ\b", body):
                has_text = True
        print(f"  {f.name}: {'HAS TEXT' if has_text else 'image-only OK'}")
        bad += has_text
    print("== Demo answers ==")
    for a in gt["demo_answers"]["applicants"]:
        print(f"  {a['name']}: {a['file']} payroll on page(s) {a['payroll_experience_pages']} of {a['total_pages']}")
    print("  near misses:", gt["demo_answers"]["near_misses"])
    if bad:
        sys.exit(f"{bad} verification failures")
    print("OK")


if __name__ == "__main__":
    main()
